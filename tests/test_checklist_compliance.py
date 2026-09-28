"""The remediation checklists, enforced and audited.

docs/Checklist for Document Remediation1.docx and
docs/ChecklistBeforeSubmittingDoc.xlsx: the text rules
(src/docx/remediation_text.py), the DOCX projection that applies them
(src/docx/docx_generator.py), the auditor that checks a finished DOCX
(src/validation/checklist_audit.py), automatic alt text
(src/images/auto_alt_text.py), the prose-table gate
(src/tables/table_extractor.py) and Docling's heading labels on OCR pages.
"""

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as open_docx
from docx.oxml.ns import qn
from PIL import Image as PILImage
from PIL import ImageDraw

from src.docx.docx_generator import (
    _HeadingLevels,
    _Segment,
    _continues,
    _describe_table,
    _stitch_prose,
    generate_docx,
)
from src.docx.remediation_text import (
    format_numbered_heading,
    has_ligatures,
    normalize_punctuation_spacing,
    remediate_prose,
    space_initialisms,
    space_units,
    unspaced_initialisms,
)
from src.models.contracts import Document, Metadata
from src.validation.checklist_audit import Status, audit_docx

WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"


def _document() -> Document:
    return Document(source_pdf_path="missing.pdf", metadata=Metadata(filename="missing.pdf"))


def _results(report) -> dict:
    return {r.item_id: r for r in report.results}


# --- text rules ------------------------------------------------------------------


class TestPunctuationSpacing:
    @pytest.mark.parametrize(
        "raw, expected",
        [
            ("apples ,pears", "apples, pears"),
            ("one,two", "one, two"),
            ("1,000 people", "1,000 people"),
            ("Smith(2005)said", "Smith (2005) said"),
            ("teacher(s) are", "teacher(s) are"),
            ("word -word", "word - word"),
            ("well-known", "well-known"),
            ("two  spaces here", "two spaces here"),
        ],
    )
    def test_checklist_spacing(self, raw: str, expected: str) -> None:
        assert normalize_punctuation_spacing(raw) == expected

    def test_ligatures_become_letters(self) -> None:
        assert has_ligatures("ﬁnd")
        assert remediate_prose("ﬁnd the eﬀect") == "find the effect"

    def test_mismapped_equals_is_repaired_but_fractions_are_not(self) -> None:
        assert remediate_prose("(r ¼ .66, p < .01)") == "(r = .66, p < .01)"
        assert remediate_prose("1¼ cups") == "1¼ cups"

    def test_urls_are_left_alone(self) -> None:
        assert remediate_prose("see www.example.com/a,b(c) now") == "see www.example.com/a,b(c) now"


class TestUnits:
    def test_space_before_units(self) -> None:
        assert space_units("5kg and 3.5km") == "5 kg and 3.5 km"

    def test_single_letter_suffixes_are_not_units(self) -> None:
        assert space_units("the 1990s, Section 2A, 3D") == "the 1990s, Section 2A, 3D"


class TestInitialisms:
    def test_letter_by_letter_initialisms_are_spaced(self) -> None:
        assert space_initialisms("The US and UK signed.") == "The U S and U K signed."

    def test_pronounceable_acronyms_are_kept(self) -> None:
        assert space_initialisms("UNESCO and NASA agreed.") == "UNESCO and NASA agreed."

    def test_hard_clusters_are_spelled(self) -> None:
        assert space_initialisms("NCERT and OECD reports") == "N C E R T and O E C D reports"

    def test_emphasis_and_roman_numerals_are_words(self) -> None:
        assert space_initialisms("Do NOT skip Part II") == "Do NOT skip Part II"

    def test_all_caps_lines_are_not_prose(self) -> None:
        assert space_initialisms("INTRODUCTION TO US POLICY") == "INTRODUCTION TO US POLICY"
        assert unspaced_initialisms("INTRODUCTION TO US POLICY") == []

    def test_idempotent(self) -> None:
        once = remediate_prose("The US ,UK(2001)said 5kg")
        assert remediate_prose(once) == once


class TestNumberedHeadings:
    @pytest.mark.parametrize(
        "raw, expected",
        [
            ("6.1 INTRODUCTION", "6.1 - INTRODUCTION"),
            ("6.1. Introduction", "6.1 - Introduction"),
            ("6.1 – Summary", "6.1 - Summary"),
            ("1.1 Scope", "1.1 - Scope"),
            ("6.1 - Done", "6.1 - Done"),
            ("2000 Years of Schooling", "2000 Years of Schooling"),
            ("UNIT 6 - RAY OPTICS", "UNIT 6 - RAY OPTICS"),
        ],
    )
    def test_checklist_hyphen_form(self, raw: str, expected: str) -> None:
        assert format_numbered_heading(remediate_prose(raw)) == expected


# --- heading hierarchy and sentence stitching ----------------------------------------


class TestHeadingLevels:
    def test_detected_levels_become_consecutive(self) -> None:
        levels = _HeadingLevels([3, 4, 3], title_is_h1=False)
        assert [levels.content(3), levels.content(4), levels.content(3)] == [1, 2, 1]

    def test_title_takes_level_one(self) -> None:
        levels = _HeadingLevels([1, 2], title_is_h1=True)
        assert levels.title() == 1
        assert [levels.content(1), levels.content(2)] == [2, 3]

    def test_no_heading_skips_a_level(self) -> None:
        levels = _HeadingLevels([1, 2, 3], title_is_h1=False)
        assert [levels.content(1), levels.content(3)] == [1, 2]

    def test_level_six_is_never_a_content_level(self) -> None:
        levels = _HeadingLevels([1, 2, 3, 4, 5, 6], title_is_h1=False)
        assert max(levels.content(level) for level in range(1, 7)) <= 5


def _para(text: str):
    return SimpleNamespace(text=text)


class TestSentenceStitching:
    def test_continues_needs_an_open_sentence(self) -> None:
        assert _continues("the reason we look at the", "evidence here.")
        assert not _continues("A full sentence ends here.", "and this one")
        assert not _continues("Short label", "lowercase start")

    def test_page_end_sentence_is_completed_from_the_next_page(self) -> None:
        first = _para("Teachers across the country reported that the")
        second = _para("curriculum was too long. Then a new idea began.")
        stitched = _stitch_prose(
            {1: [("paragraph", first)], 2: [("paragraph", second)]}, [1, 2], set()
        )
        page_one = stitched[1][0][1]
        assert [s.text for s in page_one] == [first.text, "curriculum was too long."]
        assert [s.text for s in stitched[2][0][1]] == ["Then a new idea began."]

    def test_page_ending_in_a_table_hands_nothing_on(self) -> None:
        first = _para("Teachers across the country reported that the")
        second = _para("curriculum was too long.")
        stitched = _stitch_prose(
            {1: [("paragraph", first)], 2: [("paragraph", second)]}, [1, 2], {1}
        )
        assert len(stitched[1][0][1]) == 1 and len(stitched[2]) == 1

    def test_split_paragraph_on_one_page_is_rejoined(self) -> None:
        first = _para("The study examined how teachers in three states")
        second = _para("understood their own practice.")
        stitched = _stitch_prose({1: [("paragraph", first), ("paragraph", second)]}, [1], set())
        assert len(stitched[1]) == 1 and len(stitched[1][0][1]) == 2

    def test_segment_is_a_slice_of_the_model_text(self) -> None:
        paragraph = _para("abc def")
        assert _Segment(paragraph, 4, 7).text == "def"


def test_table_summary_states_only_what_the_table_says() -> None:
    grid = [["Name", "Score"], ["A", "x"]]
    assert _describe_table(grid, {0}) == "Table with 2 rows and 2 columns. Column headings: Name; Score."


def test_table_summary_explains_its_statistics() -> None:
    grid = [["State", "Agree", "Year"], ["Bihar", "12%", "2010"], ["Kerala", "45%", "n/a"], ["Goa", "30%", "2012"]]
    assert _describe_table(grid, {0}) == (
        "Table with 4 rows and 3 columns. Column headings: State; Agree; Year. "
        "Agree: highest 45% (Kerala), lowest 12% (Bihar)."
    )


# --- the generated DOCX ---------------------------------------------------------------

_MARKDOWN = "\n\n".join(
    [
        "###### 1",
        "## 1.1 Introduction",
        "The US ,and the UK(2001)said 5kg. See www.example.com for more.",
        "## 1.2 Method",
        "Body text.",
        "<!-- pagebreak -->",
        "###### 2",
        "#### Deep heading",
        "More text.",
    ]
)


@pytest.fixture
def generated(tmp_path: Path) -> Path:
    return generate_docx(_document(), _MARKDOWN, output_path=tmp_path / "out.docx")


class TestGeneratedDocx:
    def test_heading_styles_are_times_new_roman_black_bold_not_italic(self, generated: Path) -> None:
        doc = open_docx(str(generated))
        for level in range(1, 7):
            font = doc.styles[f"Heading {level}"].font
            assert (font.name, font.bold, font.italic, str(font.color.rgb)) == (
                "Times New Roman", True, False, "000000"
            )

    def test_no_theme_fonts_remain_in_styles(self, generated: Path) -> None:
        styles = open_docx(str(generated)).styles.element
        assert not [el for el in styles.iter(qn("w:rFonts")) if el.get(qn("w:asciiTheme"))]

    def test_body_text_is_remediated_and_linked(self, generated: Path) -> None:
        doc = open_docx(str(generated))
        body = next(p for p in doc.paragraphs if "said" in "".join(t.text for t in p._p.iter(qn("w:t"))))
        text = "".join(t.text for t in body._p.iter(qn("w:t")))
        assert "The U S, and the U K (2001) said 5 kg." in text
        assert body._p.findall(qn("w:hyperlink"))

    def test_numbered_headings_and_hierarchy(self, generated: Path) -> None:
        doc = open_docx(str(generated))
        headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
        content = [h for h in headings if h[0] != "Heading 6"]
        assert content == [
            ("Heading 1", "1.1 - Introduction"),
            ("Heading 1", "1.2 - Method"),
            ("Heading 2", "Deep heading"),
        ]

    def test_no_table_of_contents_is_generated(self, generated: Path) -> None:
        """Reviewers asked for no generated index (2026-09-28); the audit
        leaves the TOC to the reviewer rather than failing the document."""
        doc = open_docx(str(generated))
        assert not any("TOC" in (t.text or "") for t in doc.element.body.iter(qn("w:instrText")))
        assert _results(audit_docx(generated))["SUB-TOC"].status is Status.MANUAL

    def test_properties_carry_no_tool_information(self, generated: Path) -> None:
        props = open_docx(str(generated)).core_properties
        assert "python-docx" not in (props.author + props.comments)
        assert props.title and props.subject and props.language == "en-US"

    def test_audit_passes_the_generated_document(self, generated: Path) -> None:
        results = _results(audit_docx(generated))
        for item in ("DR-02", "DR-06", "DR-HD-FMT", "DR-HD-ORDER", "DR-HD-HYPHEN", "DR-PUNCT",
                     "DR-11", "DR-UNITS", "SUB-LANG", "SUB-META-PRIV", "SUB-LINK"):
            assert results[item].status is Status.PASS, (item, results[item].evidence)


def test_decorative_image_carries_words_decorative_flag(tmp_path: Path) -> None:
    from docx import Document as new_docx

    from src.docx.docx_generator import _add_image

    image = tmp_path / "x.png"
    PILImage.new("RGB", (40, 40), "white").save(image)
    doc = new_docx()
    assert _add_image(doc, str(image), decorative=True)
    doc_pr = doc.element.body.find(f".//{WP}docPr")
    flags = [el for el in doc_pr.iter() if el.tag.endswith("}decorative")]
    assert flags and flags[0].get("val") == "1"


def test_unembeddable_image_leaves_no_empty_paragraph(tmp_path: Path) -> None:
    from docx import Document as new_docx

    from src.docx.docx_generator import _add_image

    bogus = tmp_path / "bogus.png"
    bogus.write_bytes(b"not an image")
    doc = new_docx()
    before = len(doc.paragraphs)
    assert not _add_image(doc, str(bogus))
    assert len(doc.paragraphs) == before


# --- the auditor on a document that breaks the rules ------------------------------------


def test_audit_fails_a_plain_python_docx_document(tmp_path: Path) -> None:
    from docx import Document as new_docx

    doc = new_docx()
    doc.add_heading("Chapter", level=3)
    doc.add_paragraph("The US ,and more text here with www.example.com inside.")
    path = tmp_path / "bad.docx"
    doc.save(str(path))
    results = _results(audit_docx(path))
    for item in ("DR-02", "DR-HD-FMT", "DR-HD-ORDER", "DR-11", "DR-PUNCT", "SUB-LINK", "SUB-META-PRIV"):
        assert results[item].status is Status.FAIL, item
    assert results["SUB-KEYBOARD"].status is Status.MANUAL


# --- automatic alt text -----------------------------------------------------------------


def _figure_image(path: Path) -> Path:
    image = PILImage.new("RGB", (200, 120), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle([20, 20, 180, 100], outline="black", width=3)
    draw.line([20, 60, 180, 60], fill="black", width=3)
    image.save(path)
    return path


def _gradient_image(path: Path) -> Path:
    image = PILImage.new("L", (200, 380))
    image.putdata([160 + y // 8 for y in range(380) for _ in range(200)])
    image.save(path)
    return path


def _images_document(tmp_path: Path) -> Document:
    from src.models.figure import AltTextStatus, Figure
    from src.models.image import Image

    document = _document()
    document.images = [
        Image(image_id=f"img-{index}", page_number=1, file_path=str(path), figure=Figure(
            alt_text="placeholder: description pending human review",
            alt_text_status=AltTextStatus.PENDING_REVIEW,
        ))
        for index, path in enumerate(
            (_figure_image(tmp_path / "figure.png"), _gradient_image(tmp_path / "blank.png"))
        )
    ]
    return document


def test_visually_empty_images_become_decorative_without_a_model(tmp_path: Path, monkeypatch) -> None:
    from src.images.auto_alt_text import apply_automatic_alt_text, is_visually_empty
    from src.models.figure import AltTextStatus

    monkeypatch.setenv("RAWRS_AI_STUB", "1")
    document = _images_document(tmp_path)
    assert is_visually_empty(document.images[1].file_path)
    assert not is_visually_empty(document.images[0].file_path)

    counts = apply_automatic_alt_text(document, nearby_text=lambda image: [])

    assert counts == {"decorative": 1, "described": 0, "unchanged": 1}
    assert document.images[1].figure.alt_text_status is AltTextStatus.DECORATIVE
    # The stub's canned text never reaches a document.
    assert document.images[0].figure.alt_text_status is AltTextStatus.PENDING_REVIEW


def test_a_real_provider_writes_speakable_alt_text(tmp_path: Path, monkeypatch) -> None:
    import src.ai.alt_text_generator as generator
    import src.ai.registry as registry
    from src.ai.alt_text_generator import AltTextResult
    from src.images.auto_alt_text import apply_automatic_alt_text
    from src.models.figure import AltTextStatus

    monkeypatch.delenv("RAWRS_AI_STUB", raising=False)
    monkeypatch.setattr(registry, "get_provider", lambda: SimpleNamespace(name="Ollama (test)"))
    monkeypatch.setattr(
        generator,
        "generate_alt_text",
        lambda request: AltTextResult(
            description="A **flowchart** of teacher beliefs.", purpose="p",
            visible_text="Beliefs", confidence=0.9,
        ),
    )
    document = _images_document(tmp_path)
    apply_automatic_alt_text(document, nearby_text=lambda image: [])

    figure = document.images[0].figure
    assert figure.alt_text_status is AltTextStatus.AI_GENERATED
    assert figure.alt_text == "A flowchart of teacher beliefs. Text in the image: Beliefs"


def test_ollama_provider_reports_unreachable_server(monkeypatch) -> None:
    from src.ai.providers.ollama import OllamaProvider

    monkeypatch.setenv("RAWRS_OLLAMA_URL", "http://127.0.0.1:9")
    capability = OllamaProvider().capabilities()
    assert not capability.available and "not reachable" in capability.unavailable_reason


# --- tables and OCR headings -------------------------------------------------------------


class TestProseTableGate:
    def test_two_column_prose_is_not_a_table(self) -> None:
        from src.tables.table_extractor import _is_running_prose

        rows = [
            ["the characterization of students and", "logical concerns, the kinds of data"],
            ["teachers, methodological choices and", "sought and their mode of treatment, all"],
            ["the kinds of questions researchers ask", "are influenced by the viewpoint held"],
        ]
        assert _is_running_prose(rows)

    def test_a_data_table_with_wrapped_headers_is_a_table(self) -> None:
        from src.tables.table_extractor import _is_running_prose

        rows = [["Low-LCE", "Mid-LCE", "High-LCE"], ["pedagogy", "pedagogy", "pedagogy"],
                ["12", "34", "9"], ["Bihar", "Kerala", "Maharashtra"]]
        assert not _is_running_prose(rows)


def test_docling_title_and_section_headers_become_heading_levels() -> None:
    from src.ocr.docling_engine import _layout_heading_levels

    def item(label, text, level=None):
        return SimpleNamespace(label=SimpleNamespace(value=label), text=text, level=level)

    docling_document = SimpleNamespace(iterate_items=lambda: iter([
        (item("title", "Developing Your Research Question"), 0),
        (item("section_header", "THE IMPORTANCE OF GOOD QUESTIONS", 1), 1),
        (item("text", "Body text."), 1),
        (item("section_header", "Not a line of the page", 1), 1),
    ]))
    cleaned = "Developing Your Research Question\nTHE IMPORTANCE OF GOOD QUESTIONS\nBody text."
    assert _layout_heading_levels(docling_document, cleaned) == {
        "Developing Your Research Question": 1,
        "THE IMPORTANCE OF GOOD QUESTIONS": 2,
    }


def test_ocr_heading_levels_reach_detected_headings() -> None:
    from src.headings.heading_detector import detect_headings
    from src.models.contracts import ExtractionMethod, Page

    document = _document()
    document.pages = [Page(
        page_number=1,
        cleaned_text="Developing Your Research Question\nSome body text follows here.",
        extraction_method=ExtractionMethod.DOCLING,
        ocr_heading_levels={"Developing Your Research Question": 1},
    )]
    headings = [h for h in detect_headings(document).headings if not h.is_page_marker]
    assert [(h.text, h.level.value, h.source) for h in headings] == [
        ("Developing Your Research Question", 1, "docling_layout")
    ]


def test_repeated_icons_and_ornaments_are_decorative(tmp_path: Path) -> None:
    from src.images.auto_alt_text import is_ornament, recurring_image_ids
    from src.models.image import Image

    icon = _figure_image(tmp_path / "icon.png")
    figure = tmp_path / "figure.png"
    drawn = PILImage.new("RGB", (200, 120), "white")
    ImageDraw.Draw(drawn).ellipse([10, 10, 190, 110], outline="black", width=4)
    drawn.save(figure)
    strip = tmp_path / "strip.png"
    PILImage.new("RGB", (300, 20), "gray").save(strip)

    images = [
        Image(image_id="icon-p1", page_number=1, file_path=str(icon)),
        Image(image_id="icon-p4", page_number=4, file_path=str(icon)),
        Image(image_id="figure", page_number=2, file_path=str(figure)),
    ]
    assert recurring_image_ids(images) == {"icon-p1", "icon-p4"}
    assert is_ornament(str(strip)) and not is_ornament(str(figure))


# --- the API ------------------------------------------------------------------------------


def test_checklist_endpoint_audits_the_current_export(tmp_path: Path, monkeypatch) -> None:
    from fastapi.testclient import TestClient

    from src.api import document_store, jobs
    from src.api.jobs import Job, JobStatus
    from src.api.main import app
    from src.models.contracts import Heading, HeadingLevel, Page
    from src.models.document import ProcessingStatus
    from src.pipeline.phase1_pipeline import PipelineResult

    monkeypatch.setattr(jobs, "JOBS_DIR", tmp_path / "jobs")
    monkeypatch.setattr(jobs, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(document_store, "DOCUMENTS_DIR", tmp_path / "documents")
    document = _document()
    document.pages = [Page(page_number=1)]
    document.headings = [
        Heading(text="Introduction", level=HeadingLevel.H1, page_number=1, document_order=0,
                source="mathpix", source_line=1),
    ]
    docx_path = tmp_path / "chk.docx"
    job = Job(
        job_id="chk1", filename="test.pdf", pdf_path=Path("uploads/chk1.pdf"),
        status=JobStatus.COMPLETE,
        result=PipelineResult(
            source_pdf_path="missing.pdf", success=True, status=ProcessingStatus.VALIDATED,
            duration_seconds=1.0, document=document, docx_path=docx_path,
            docx_generated_at_version=-1,
        ),
    )
    with jobs._lock:
        jobs._jobs["chk1"] = job
    try:
        response = TestClient(app).get("/api/documents/chk1/checklist")
    finally:
        with jobs._lock:
            jobs._jobs.clear()

    assert response.status_code == 200
    body = response.json()
    assert docx_path.is_file()  # the export was refreshed before it was audited
    ids = {r["item_id"] for r in body["results"]}
    assert {"DR-02", "DR-06", "SUB-TOC", "SUB-META", "SUB-KEYBOARD"} <= ids
    assert set(body["summary"]) == {"pass", "fail", "warn", "not_applicable", "manual"}


def test_charts_with_numbers_get_a_plain_language_explanation(tmp_path: Path, monkeypatch) -> None:
    import src.ai.alt_text_generator as generator
    import src.ai.registry as registry
    from src.ai.alt_text_generator import AltTextResult
    from src.images.auto_alt_text import apply_automatic_alt_text, needs_explanation

    assert needs_explanation("CHART", "45% agree and 22% disagree.")
    assert not needs_explanation("DIAGRAM", "Four boxes, 3 arrows.")
    assert not needs_explanation("CHART", "45% agree. This means most agree.")

    provider = SimpleNamespace(
        name="Ollama (test)",
        explain_statistics=lambda description: "This means about twice as many agree as disagree.",
    )
    monkeypatch.delenv("RAWRS_AI_STUB", raising=False)
    monkeypatch.setattr(registry, "get_provider", lambda: provider)
    monkeypatch.setattr(
        generator,
        "generate_alt_text",
        lambda request: AltTextResult(
            description="A bar chart: 45% agree, 22% disagree.", purpose="p",
            visible_text="None", confidence=0.9, image_type="CHART",
        ),
    )
    document = _images_document(tmp_path)
    apply_automatic_alt_text(document, nearby_text=lambda image: [])
    assert document.images[0].figure.alt_text == (
        "A bar chart: 45% agree, 22% disagree. This means about twice as many agree as disagree."
    )


def test_a_note_reference_without_a_body_is_dropped_not_left_dangling(tmp_path: Path) -> None:
    markdown = "###### 1\n\nText with a stray marker[^ghost] and a real one[^real].\n\n[^real]: The real note."
    path = generate_docx(_document(), markdown, output_path=tmp_path / "notes.docx")
    doc = open_docx(str(path))
    ids = [el.get(qn("w:id")) for el in doc.element.body.iter(qn("w:footnoteReference"))]
    assert ids == ["2"]  # "ghost" claimed id 1 first, has no body, and is gone
    assert _results(audit_docx(path))["SUB-NOTES"].status is Status.PASS
