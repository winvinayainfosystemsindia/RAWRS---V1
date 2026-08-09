"""W-2b — a prose edit is a semantic decision.

The first actual prose-editing capability in RAWRS. Every other semantic
object already had one — a heading's level and text, a note's body, a
table's cells, an image's alt text — while a typo in an ordinary sentence
had no edit path at all.

What these tests pin is that it arrives as an ordinary correction rather
than as a new mechanism: one request is one ``CorrectionRecord`` in one
transaction, the mutation is performed by ``ParagraphVerifier.apply()``
through ``engine.apply_correction()``, undo is
``SemanticVerifier.revert()``'s generic swap with no paragraph-specific
logic, and both projections change only because they read
``Paragraph.text`` from the Semantic Document. Nothing parses Markdown or
DOCX in either direction.

The target is always ``Paragraph.id`` (W-2a), never a list index, a page
position or the text itself — which is what keeps two paragraphs carrying
identical prose independently editable.
"""

import time
from pathlib import Path
from typing import List, Optional

import pytest
from fastapi.testclient import TestClient

from src.api.main import app
from src.docx.docx_generator import generate_docx
from src.markdown.markdown_builder import build_markdown
from src.models.bounding_box import BoundingBox
from src.models.contracts import Document, Metadata, Page, Paragraph, TextBlock
from src.models.correction import APPLIED_STATUSES, CorrectionRecord, CorrectionStatus
from src.verification.engine import NonLifoUndoError, engine
from src.verification.paragraphs import ASSET_TYPE, TEXT_EDIT, ParagraphVerifier

client = TestClient(app)

SAMPLE_PDF = (
    Path(__file__).resolve().parents[1]
    / "samples"
    / "benchmark"
    / "pdfs"
    / "5.Teachingas a profession_Calderhead.pdf"
)


# ── In-memory fixtures: fast, isolated, no pipeline ──────────────────────


def _document() -> Document:
    """Three paragraphs, two of which say exactly the same thing.

    The duplicate pair is not decoration: it is the case a text-keyed or
    position-keyed edit rail would get wrong, and the reason identity is
    derived from the source block.
    """
    blocks = [
        TextBlock(
            page_number=1,
            text=f"Line {i}",
            bbox=BoundingBox(x0=72, y0=40 + i * 20, x1=400, y1=56 + i * 20),
            order=i,
        )
        for i in range(3)
    ]
    paragraphs = [
        Paragraph(page_number=1, text="Ibid.", source_block_ids=[blocks[0].block_id]),
        Paragraph(page_number=1, text="Ibid.", source_block_ids=[blocks[1].block_id]),
        Paragraph(page_number=1, text="Original prose.", source_block_ids=[blocks[2].block_id]),
    ]
    return Document(
        source_pdf_path="x.pdf",
        metadata=Metadata(filename="x.pdf", page_count=1),
        pages=[Page(page_number=1, cleaned_text="Prose.")],
        blocks=blocks,
        paragraphs=paragraphs,
    )


def _edit(
    document: Document,
    paragraph: Paragraph,
    new_text: str,
    transaction_id: Optional[str] = None,
) -> CorrectionRecord:
    """One prose edit, exactly as the endpoint records it."""
    correction = CorrectionRecord(
        object_type=ASSET_TYPE,
        object_id=paragraph.id,
        field=TEXT_EDIT,
        original_value=paragraph.text,
        proposed_value=new_text,
        reason="Reviewer edited the paragraph text.",
        reason_code="PARAGRAPH_TEXT_EDITED_BY_REVIEWER",
        provider="manual_reviewer",
        status=CorrectionStatus.EDITED,
        transaction_id=transaction_id or f"txn-{len(document.corrections)}",
    )
    document.corrections.append(correction)
    engine.apply_correction(document, correction)
    return correction


class TestTheVerifierIsRegistered:
    def test_paragraph_is_a_registered_asset_type(self) -> None:
        assert isinstance(engine._verifiers[ASSET_TYPE], ParagraphVerifier)  # noqa: SLF001

    def test_it_owns_one_rule_id_for_one_editable_field(self) -> None:
        table = ParagraphVerifier().rule_table()
        assert list(table) == [TEXT_EDIT]
        assert table[TEXT_EDIT].rule_id == "PARAGRAPH_001"
        assert table[TEXT_EDIT].reason_code == "PARAGRAPH_TEXT_EDITED_BY_REVIEWER"

    def test_it_proposes_nothing_of_its_own(self) -> None:
        # RAWRS has no basis for claiming a sentence should read
        # differently, so this verifier is a carrier, not a producer.
        assert ParagraphVerifier().inspect(_document()) == []


class TestTheEditIsAppliedByTheRail:
    def test_the_correction_changes_the_paragraph_text(self) -> None:
        document = _document()
        target = document.paragraphs[2]

        _edit(document, target, "Corrected prose.")

        assert target.text == "Corrected prose."

    def test_the_document_version_is_bumped_so_exports_invalidate(self) -> None:
        document = _document()
        before = document.version

        _edit(document, document.paragraphs[2], "Corrected prose.")

        assert document.version == before + 1

    def test_the_correction_carries_the_edit_it_made(self) -> None:
        document = _document()
        correction = _edit(document, document.paragraphs[2], "Corrected prose.")

        assert len(document.corrections) == 1
        assert correction.object_type == "paragraph"
        assert correction.object_id == document.paragraphs[2].id
        assert correction.field == "text_edit"
        assert correction.original_value == "Original prose."
        assert correction.proposed_value == "Corrected prose."
        assert correction.status is CorrectionStatus.EDITED
        assert correction.status in APPLIED_STATUSES
        assert correction.provider == "manual_reviewer"
        assert correction.transaction_id is not None

    def test_applying_twice_writes_the_same_text(self) -> None:
        document = _document()
        target = document.paragraphs[2]
        correction = _edit(document, target, "Corrected prose.")

        engine.apply_correction(document, correction)

        assert target.text == "Corrected prose."

    def test_source_blocks_are_never_touched(self) -> None:
        # The edit changes what the paragraph says, not what extraction
        # found — the split that keeps inline formatting and footnote
        # marker placement working off the original lines.
        document = _document()
        before = [b.text for b in document.blocks]

        _edit(document, document.paragraphs[2], "Corrected prose.")

        assert [b.text for b in document.blocks] == before

    def test_no_other_paragraph_moves_or_changes(self) -> None:
        document = _document()
        ids_before = [p.id for p in document.paragraphs]

        _edit(document, document.paragraphs[2], "Corrected prose.")

        assert [p.id for p in document.paragraphs] == ids_before
        assert document.paragraphs[0].text == "Ibid."
        assert document.paragraphs[1].text == "Ibid."


class TestIdentityIsTheOnlyTarget:
    def test_paragraphs_with_identical_text_are_edited_independently(self) -> None:
        document = _document()
        first, second = document.paragraphs[0], document.paragraphs[1]

        _edit(document, first, "Edited the first one.")

        assert first.text == "Edited the first one."
        assert second.text == "Ibid."

    def test_a_correction_naming_no_paragraph_is_a_clean_no_op(self) -> None:
        document = _document()
        texts = [p.text for p in document.paragraphs]

        correction = CorrectionRecord(
            object_type=ASSET_TYPE,
            object_id="paragraph-does-not-exist",
            field=TEXT_EDIT,
            original_value="",
            proposed_value="Should land nowhere.",
            status=CorrectionStatus.EDITED,
        )
        engine.apply_correction(document, correction)  # must not raise

        assert [p.text for p in document.paragraphs] == texts

    def test_a_correction_for_another_field_is_ignored(self) -> None:
        # object_type routes to this verifier; the field is what apply()
        # dispatches on, so a field it does not own must change nothing.
        document = _document()
        texts = [p.text for p in document.paragraphs]

        ParagraphVerifier().apply(
            document,
            CorrectionRecord(
                object_type=ASSET_TYPE,
                object_id=document.paragraphs[2].id,
                field="some_other_field",
                original_value="Original prose.",
                proposed_value="Should land nowhere.",
            ),
        )

        assert [p.text for p in document.paragraphs] == texts


class TestUndoIsTheGenericOne:
    def test_undo_restores_the_previous_wording(self) -> None:
        document = _document()
        target = document.paragraphs[2]
        _edit(document, target, "Corrected prose.")

        reverted = engine.undo_last_transaction(document)

        assert target.text == "Original prose."
        assert [c.status for c in reverted] == [CorrectionStatus.REVERTED]

    def test_undo_restores_the_previous_decision_not_the_original(self) -> None:
        document = _document()
        target = document.paragraphs[2]
        _edit(document, target, "First edit.")
        _edit(document, target, "Second edit.")

        engine.undo_last_transaction(document)

        assert target.text == "First edit."

    def test_the_history_is_append_only(self) -> None:
        document = _document()
        _edit(document, document.paragraphs[2], "Corrected prose.")

        engine.undo_last_transaction(document)

        assert len(document.corrections) == 1  # moved to REVERTED, not deleted
        assert document.corrections[0].status is CorrectionStatus.REVERTED

    def test_out_of_order_undo_is_refused(self) -> None:
        # LIFO protection: reverting an older edit would replay a swap
        # whose "original" is no longer what the paragraph says, silently
        # clobbering the later decision.
        document = _document()
        older = _edit(document, document.paragraphs[2], "First edit.")
        _edit(document, document.paragraphs[2], "Second edit.")

        with pytest.raises(NonLifoUndoError):
            engine.undo_last_transaction(document, expect_transaction_id=older.transaction_id)

        assert document.paragraphs[2].text == "Second edit."

    def test_two_paragraphs_edited_separately_undo_separately(self) -> None:
        document = _document()
        first, third = document.paragraphs[0], document.paragraphs[2]
        _edit(document, first, "Edited the first one.")
        _edit(document, third, "Edited the third one.")

        engine.undo_last_transaction(document)

        assert third.text == "Original prose."
        assert first.text == "Edited the first one."  # untouched by the undo


# ── The real document: API surface and both projections ──────────────────


@pytest.fixture(scope="module")
def job_id() -> str:
    """One corpus PDF through the whole pipeline.

    Processing is asynchronous and _require_document() answers 409 until
    it finishes, so every request below acts on a fully-built Semantic
    Document.
    """
    with SAMPLE_PDF.open("rb") as handle:
        response = client.post(
            "/api/documents", files={"file": (SAMPLE_PDF.name, handle, "application/pdf")}
        )
    assert response.status_code in (200, 201), response.text
    job = response.json()["job_id"]

    deadline = time.time() + 300
    while time.time() < deadline:
        if client.get(f"/api/documents/{job}").json().get("status") not in ("queued", "processing"):
            break
        time.sleep(1)
    else:  # pragma: no cover - only on a pathologically slow machine
        pytest.skip("document did not finish processing in time")
    return job


def _document_for(job: str) -> Document:
    from src.api.routes import _require_document

    return _require_document(job)


@pytest.fixture(scope="module")
def targets(job_id: str) -> List[str]:
    """One substantial paragraph per mutating test, longest first.

    Chosen once, before anything is edited, and by id rather than by
    position — two reasons. A short paragraph's text is a substring of
    longer ones ("Teachers" appears inside "Teachers similarly face..."),
    so "the previous wording is gone" cannot be asserted against it; and
    the document is module-scoped, so tests must not compete for the same
    paragraph.
    """
    ordered = sorted(_document_for(job_id).paragraphs, key=lambda p: len(p.text), reverse=True)
    return [p.id for p in ordered]


def _paragraph(job: str, paragraph_id: str) -> Paragraph:
    return next(p for p in _document_for(job).paragraphs if p.id == paragraph_id)


def _paragraph_corrections(job: str) -> List[CorrectionRecord]:
    return [c for c in _document_for(job).corrections if c.object_type == "paragraph"]


class TestTheEndpoint:
    def test_paragraphs_are_listed_with_their_identities(self, job_id: str) -> None:
        listed = client.get(f"/api/documents/{job_id}/paragraphs").json()["paragraphs"]

        assert listed
        assert all(p["paragraph_id"] for p in listed)

    def test_an_edit_is_one_correction_and_changes_the_document(
        self, job_id: str, targets: List[str]
    ) -> None:
        target = _paragraph(job_id, targets[0])
        before_text, before_count = target.text, len(_paragraph_corrections(job_id))

        response = client.patch(
            f"/api/documents/{job_id}/paragraphs/{target.id}",
            json={"text": "This sentence was corrected by a reviewer."},
        )

        assert response.status_code == 200
        assert response.json()["text"] == "This sentence was corrected by a reviewer."
        assert target.text == "This sentence was corrected by a reviewer."

        corrections = _paragraph_corrections(job_id)
        assert len(corrections) == before_count + 1
        correction = corrections[-1]
        assert correction.object_id == target.id
        assert correction.field == "text_edit"
        assert correction.original_value == before_text
        assert correction.reason_code == "PARAGRAPH_TEXT_EDITED_BY_REVIEWER"
        assert correction.status in APPLIED_STATUSES

    def test_undo_through_the_generic_corrections_endpoint(
        self, job_id: str, targets: List[str]
    ) -> None:
        target = _paragraph(job_id, targets[1])
        original = target.text

        client.patch(
            f"/api/documents/{job_id}/paragraphs/{target.id}", json={"text": "Temporarily wrong."}
        )
        correction = _paragraph_corrections(job_id)[-1]

        undo = client.patch(
            f"/api/documents/{job_id}/corrections/{correction.correction_id}",
            json={"action": "undo"},
        )

        assert undo.status_code == 200
        assert target.text == original

    def test_an_unknown_paragraph_is_refused(self, job_id: str) -> None:
        response = client.patch(
            f"/api/documents/{job_id}/paragraphs/paragraph-nope", json={"text": "Anything."}
        )
        assert response.status_code == 404

    @pytest.mark.parametrize("text", ["", "   "])
    def test_blank_text_is_refused(self, job_id: str, targets: List[str], text: str) -> None:
        # Existing policy, taken from the note-body endpoint: a paragraph
        # that says nothing is a deletion, which a text edit is not.
        target = _paragraph(job_id, targets[2])
        before = target.text

        response = client.patch(
            f"/api/documents/{job_id}/paragraphs/{target.id}", json={"text": text}
        )

        assert response.status_code == 422
        assert target.text == before

    def test_an_edit_to_identical_text_is_still_recorded(
        self, job_id: str, targets: List[str]
    ) -> None:
        # No existing endpoint de-duplicates a no-change edit (metadata and
        # footnote bodies both record one), so this preserves that
        # behaviour rather than inventing a policy: the reviewer made a
        # decision, and the log says so. Harmless — the mutation is a
        # no-op and undo restores the same string.
        target = _paragraph(job_id, targets[3])
        before_count = len(_paragraph_corrections(job_id))

        client.patch(
            f"/api/documents/{job_id}/paragraphs/{target.id}", json={"text": target.text}
        )

        corrections = _paragraph_corrections(job_id)
        assert len(corrections) == before_count + 1
        assert corrections[-1].original_value == corrections[-1].proposed_value


class TestBothProjectionsFollowTheModel:
    """The projections are observers of the semantic mutation.

    Neither renderer knows a paragraph edit happened: Markdown changes
    because ``_render_paragraph`` reads ``Paragraph.text``, and the DOCX
    changes because it is generated from that Markdown. Nothing parses
    either format back into the model.
    """

    # One phrase per test: the document is module-scoped, so a shared
    # marker would still be present from the earlier edit.
    PROJECTED = "Zarquon was here, said the reviewer."
    REVERTED = "Slartibartfast was here instead."

    def test_markdown_and_docx_both_carry_the_edited_text(
        self, job_id: str, targets: List[str], tmp_path
    ) -> None:
        from docx import Document as DocxDocument

        document = _document_for(job_id)
        target = _paragraph(job_id, targets[4])
        original = target.text

        assert original in build_markdown(document)

        client.patch(
            f"/api/documents/{job_id}/paragraphs/{target.id}", json={"text": self.PROJECTED}
        )

        markdown = build_markdown(document)
        assert self.PROJECTED in markdown
        assert original not in markdown

        docx_path = generate_docx(document, markdown, tmp_path / "edited.docx")
        docx_text = "\n".join(p.text for p in DocxDocument(str(docx_path)).paragraphs)
        assert self.PROJECTED in docx_text
        assert original not in docx_text

    def test_undo_returns_the_projection_to_the_original(
        self, job_id: str, targets: List[str]
    ) -> None:
        document = _document_for(job_id)
        target = _paragraph(job_id, targets[5])
        original = target.text

        client.patch(
            f"/api/documents/{job_id}/paragraphs/{target.id}", json={"text": self.REVERTED}
        )
        assert self.REVERTED in build_markdown(document)

        correction = _paragraph_corrections(job_id)[-1]
        client.patch(
            f"/api/documents/{job_id}/corrections/{correction.correction_id}",
            json={"action": "undo"},
        )

        markdown = build_markdown(document)
        assert original in markdown
        assert self.REVERTED not in markdown


class TestAnUneditedDocumentIsUnchanged:
    """Registering an eleventh asset type must cost documents nothing.

    The real regression risk in adding a verifier is that it starts
    contributing findings — and therefore corrections, validation issues
    and export blocks — to every document that runs inspection.
    """

    def test_inspection_produces_no_paragraph_findings(self) -> None:
        document = _document()

        findings = engine.run_inspection(document)

        assert [f for f in findings if f.asset_type == ASSET_TYPE] == []
        assert document.corrections == []

    def test_an_unedited_corpus_document_projects_identically(self) -> None:
        from src.ocr.extractor import extract_text
        from src.parser.pdf_parser import parse_pdf
        from src.structure.structure_detector import detect_structure

        document = detect_structure(extract_text(parse_pdf(SAMPLE_PDF)))
        before_paragraphs = [
            (p.id, p.text, p.page_number, tuple(p.source_block_ids)) for p in document.paragraphs
        ]
        before_markdown = build_markdown(document)

        engine.run_inspection(document)

        after = [
            (p.id, p.text, p.page_number, tuple(p.source_block_ids)) for p in document.paragraphs
        ]
        assert after == before_paragraphs
        assert build_markdown(document) == before_markdown
