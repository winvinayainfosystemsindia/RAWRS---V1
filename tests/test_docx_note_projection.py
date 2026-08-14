"""P4c-2 — the DOCX projection materializes notes from the traversal.

DOCX used to learn a note's body by parsing the markdown's ``[^label]: body``
line back out, and resolved that note's part, its ``w:id`` and its body through
the label — which src/models/footnote.py says in as many words is *not*
identity: the label is derived from the printed number and the body page, so a
reviewer's renumbering rewrites it while the note stays the same note.

Now a ``NOTE_DEFINITION`` node names the note and ``Footnote.note_type`` names
its part. The corpus contains zero footnotes (7 of 7 notes are endnotes), so
every footnote claim below is a fixture claim and is not made of the corpus.
"""

import zipfile
from pathlib import Path
from typing import List, Optional, Tuple
from xml.etree import ElementTree as etree

from src.benchmark.projection import check_note_projection
from src.docx.docx_generator import generate_docx
from src.markdown.markdown_builder import build_markdown
from src.models.bounding_box import BoundingBox
from src.models.contracts import Document, Metadata, Page, Paragraph, TextBlock
from src.models.footnote import Footnote, NoteType

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_FOOTNOTES = "word/footnotes.xml"
_ENDNOTES = "word/endnotes.xml"


def _note(
    note_type: NoteType,
    number: int,
    note_id: str,
    page: int = 1,
    body: Optional[str] = None,
) -> Footnote:
    """A note whose marker sits at offset 7 of its own anchor line."""
    return Footnote(
        note_type=note_type,
        number=number,
        marker=str(number),
        anchor_page_number=page,
        anchor_text=f"A claim{number} on page {page}.",
        anchor_offset=7,
        body=body or f"Body of {note_id}.",
        body_page_number=page,
        body_source_text=f"{number}. Body of {note_id}.",
        footnote_id=note_id,
    )


def _document(notes: List[Footnote]) -> Document:
    """A Document whose every note anchor line is real prose on a real page,
    so the anchor page is a stream page and the reference is emitted natively."""
    pages = sorted({n.anchor_page_number for n in notes}) or [1]
    blocks, paragraphs = [], []
    for order, note in enumerate(notes):
        block = TextBlock(
            page_number=note.anchor_page_number,
            text=note.anchor_text,
            bbox=BoundingBox(x0=72, y0=40 + order * 20, x1=400, y1=56 + order * 20),
            order=order,
        )
        blocks.append(block)
        paragraphs.append(
            Paragraph(
                page_number=note.anchor_page_number,
                text=note.anchor_text,
                source_block_ids=[block.block_id],
            )
        )
    return Document(
        source_pdf_path="notes.pdf",
        metadata=Metadata(filename="notes.pdf", page_count=len(pages)),
        pages=[
            Page(
                page_number=p,
                cleaned_text="\n".join(b.text for b in blocks if b.page_number == p),
            )
            for p in pages
        ],
        blocks=blocks,
        paragraphs=paragraphs,
        footnotes=notes,
    )


def _render(document: Document, tmp_path: Path, markdown: Optional[str] = None) -> Path:
    out = tmp_path / "notes.docx"
    generate_docx(
        document, markdown if markdown is not None else build_markdown(document), out
    )
    return out


def _entries(path: Path, part: str, tag: str) -> List[Tuple[int, str]]:
    """(id, body) for every real note in a part, in document order. Word's
    separator entries (ids -1/0) belong to no Footnote and are excluded."""
    with zipfile.ZipFile(str(path)) as zf:
        if part not in zf.namelist():
            return []
        root = etree.fromstring(zf.read(part))
    out = []
    for el in root.findall(f"{_W}{tag}"):
        if int(el.get(f"{_W}id", "-99")) < 1:
            continue
        text = "".join(t.text or "" for t in el.iter(f"{_W}t")).strip()
        out.append((int(el.get(f"{_W}id")), text))
    return out


def _references(path: Path) -> List[Tuple[str, int]]:
    with zipfile.ZipFile(str(path)) as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    return [
        (el.tag.split("}")[1], int(el.get(f"{_W}id")))
        for el in root.iter()
        if el.tag in (f"{_W}footnoteReference", f"{_W}endnoteReference")
    ]


def _body_texts(path: Path) -> List[str]:
    from docx import Document as DocxDocument

    return [p.text for p in DocxDocument(str(path)).paragraphs if p.text.strip()]


class TestFootnoteMaterialization:
    """The corpus has no footnotes at all, so these six are the whole
    evidence for the footnote path (P4C2 audit §7)."""

    def test_a_single_footnote_reaches_the_footnotes_part(self, tmp_path: Path) -> None:
        document = _document([_note(NoteType.FOOTNOTE, 1, "fn-a")])
        path = _render(document, tmp_path)

        assert _entries(path, _FOOTNOTES, "footnote") == [(1, "Body of fn-a.")]
        assert _entries(path, _ENDNOTES, "endnote") == []
        assert _references(path) == [("footnoteReference", 1)]

    def test_a_footnote_and_an_endnote_keep_their_own_parts(self, tmp_path: Path) -> None:
        document = _document(
            [_note(NoteType.FOOTNOTE, 1, "fn-a"), _note(NoteType.ENDNOTE, 2, "fn-b")]
        )
        path = _render(document, tmp_path)

        assert _entries(path, _FOOTNOTES, "footnote") == [(1, "Body of fn-a.")]
        assert _entries(path, _ENDNOTES, "endnote") == [(1, "Body of fn-b.")]

    def test_the_two_id_spaces_are_independent(self, tmp_path: Path) -> None:
        """Each part is its own w:id space, both starting at 1 — a shared
        counter would renumber one part's references into the other's."""
        document = _document(
            [
                _note(NoteType.FOOTNOTE, 1, "fn-a"),
                _note(NoteType.ENDNOTE, 2, "fn-b"),
                _note(NoteType.FOOTNOTE, 3, "fn-c"),
                _note(NoteType.ENDNOTE, 4, "fn-d"),
            ]
        )
        path = _render(document, tmp_path)

        assert [i for i, _ in _entries(path, _FOOTNOTES, "footnote")] == [1, 2]
        assert [i for i, _ in _entries(path, _ENDNOTES, "endnote")] == [1, 2]
        assert sorted(_references(path)) == [
            ("endnoteReference", 1),
            ("endnoteReference", 2),
            ("footnoteReference", 1),
            ("footnoteReference", 2),
        ]

    def test_two_notes_with_the_same_label_are_two_notes(self, tmp_path: Path) -> None:
        """``Footnote.label`` is "p{body_page}-{number}", so two notes that
        print the same number on the same page derive the *same* label. Keyed
        by label the second body overwrote the first; keyed by identity both
        survive."""
        document = _document(
            [
                _note(NoteType.FOOTNOTE, 1, "fn-a", body="First body."),
                _note(NoteType.FOOTNOTE, 1, "fn-b", body="Second body."),
            ]
        )
        assert document.footnotes[0].label == document.footnotes[1].label

        path = _render(document, tmp_path)
        assert _entries(path, _FOOTNOTES, "footnote") == [
            (1, "First body."),
            (2, "Second body."),
        ]

    def test_renumbering_after_the_markdown_was_built_keeps_the_note(
        self, tmp_path: Path
    ) -> None:
        """A reviewer's renumbering rewrites the label. Identity does not
        move, so the body still materializes exactly once."""
        note = _note(NoteType.FOOTNOTE, 1, "fn-a")
        document = _document([note])
        stale_markdown = build_markdown(document)

        note.number = 4  # the label is now p1-4; the note is still fn-a

        path = _render(document, tmp_path, markdown=stale_markdown)
        assert _entries(path, _FOOTNOTES, "footnote") == [(1, "Body of fn-a.")]

    def test_a_note_whose_reference_cannot_be_located_still_materializes(
        self, tmp_path: Path
    ) -> None:
        """The body is the note's, not the prose's. An unlocatable marker
        costs the reference, never the body — and invents nothing."""
        note = _note(NoteType.FOOTNOTE, 1, "fn-a")
        note.anchor_offset = None
        note.marker = "††"  # a marker no line contains
        document = _document([note])
        path = _render(document, tmp_path)

        assert _entries(path, _FOOTNOTES, "footnote") == [(1, "Body of fn-a.")]
        assert _references(path) == []


class TestModelBeatsMarkdown:
    def test_a_reviewer_body_edit_reaches_the_package(self, tmp_path: Path) -> None:
        note = _note(NoteType.ENDNOTE, 1, "fn-a", body="Detected body.")
        document = _document([note])
        stale_markdown = build_markdown(document)

        note.body = "Reviewer body."  # the correction rail's effect

        path = _render(document, tmp_path, markdown=stale_markdown)
        assert _entries(path, _ENDNOTES, "endnote") == [(1, "Reviewer body.")]
        assert "Detected body." in stale_markdown

    def test_markdown_without_the_definition_line_still_materializes(
        self, tmp_path: Path
    ) -> None:
        """The definition line was the transport. Removing it used to lose the
        body entirely; the node carries it now."""
        document = _document([_note(NoteType.ENDNOTE, 1, "fn-a")])
        stripped = "\n".join(
            line
            for line in build_markdown(document).splitlines()
            if not line.strip().startswith("[^")
        )
        path = _render(document, tmp_path, markdown=stripped)
        assert _entries(path, _ENDNOTES, "endnote") == [(1, "Body of fn-a.")]

    def test_the_definition_line_does_not_render_as_prose(self, tmp_path: Path) -> None:
        document = _document([_note(NoteType.ENDNOTE, 1, "fn-a")])
        path = _render(document, tmp_path)
        assert not any(text.startswith("[^") for text in _body_texts(path))
        assert not any("Body of fn-a." in text for text in _body_texts(path))


class TestFallbackPreserved:
    def test_hand_written_markdown_still_renders_its_notes(self, tmp_path: Path) -> None:
        """A Document with no traversal has only its markdown, so the
        definition line remains the note's only record."""
        document = Document(source_pdf_path="x.pdf", metadata=Metadata(filename="x.pdf"))
        path = tmp_path / "hand.docx"
        generate_docx(document, "A claim[^p1-1].\n\n[^p1-1]: Hand-written body.", path)

        assert _entries(path, _FOOTNOTES, "footnote") == [(1, "Hand-written body.")]
        assert _references(path) == [("footnoteReference", 1)]


class TestNoteProjectionInvariant:
    """PI-12, run over the package the renderer actually produced."""

    def test_a_mixed_document_satisfies_pi12(self, tmp_path: Path) -> None:
        document = _document(
            [
                _note(NoteType.FOOTNOTE, 1, "fn-a"),
                _note(NoteType.ENDNOTE, 2, "fn-b"),
                _note(NoteType.FOOTNOTE, 3, "fn-c", page=2),
            ]
        )
        path = _render(document, tmp_path)

        report = check_note_projection(document, path, name="mixed")
        assert report.violations == []
        assert report.counts == {
            "notes": 3,
            "note_nodes": 3,
            "footnote_bodies": 2,
            "endnote_bodies": 1,
            "references": 3,
        }

    def test_pi12_catches_a_body_that_reached_no_part(self, tmp_path: Path) -> None:
        document = _document([_note(NoteType.ENDNOTE, 1, "fn-a")])
        empty = tmp_path / "empty.docx"
        generate_docx(
            Document(source_pdf_path="x.pdf", metadata=Metadata(filename="x.pdf")),
            "Prose with no notes.",
            empty,
        )

        report = check_note_projection(document, empty, name="lost")
        assert [v.kind for v in report.violations] == ["content_loss"]

    def test_pi12_catches_a_body_in_the_wrong_part(self, tmp_path: Path) -> None:
        footnote_doc = _document([_note(NoteType.FOOTNOTE, 1, "fn-a")])
        path = _render(footnote_doc, tmp_path)

        claims_endnote = _document([_note(NoteType.ENDNOTE, 1, "fn-a")])
        report = check_note_projection(claims_endnote, path, name="wrong")
        assert {v.kind for v in report.violations} == {"content_loss", "invented_object"}

    def test_pi12_is_clean_on_a_document_with_no_notes(self, tmp_path: Path) -> None:
        document = Document(source_pdf_path="x.pdf", metadata=Metadata(filename="x.pdf"))
        path = tmp_path / "none.docx"
        generate_docx(document, "Prose only.", path)

        report = check_note_projection(document, path, name="none")
        assert report.violations == []
        assert report.counts == {"notes": 0}
