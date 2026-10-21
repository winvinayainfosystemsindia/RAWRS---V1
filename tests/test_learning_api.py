"""The learning endpoints (src/api/learning_routes.py) and the download snapshot.

Same synthetic-job fixture pattern as tests/test_corrections_api.py; every
output root is redirected to tmp_path so nothing touches outputs/.
"""

import json
import uuid
from datetime import datetime, timezone

import pytest
from docx import Document as WordDocument


@pytest.fixture
def client():
    from fastapi.testclient import TestClient
    from src.api.main import app

    return TestClient(app)


@pytest.fixture
def roots(tmp_path, monkeypatch):
    monkeypatch.setattr("src.api.learning_routes.DEFAULT_OUTPUT_ROOT", tmp_path)
    monkeypatch.setattr("src.api.learning_routes.DOCUMENTS_DIR", tmp_path / "documents")
    monkeypatch.setattr("src.api.routes.DEFAULT_OUTPUT_ROOT", tmp_path)
    monkeypatch.setenv("RAWRS_LEARNING_SNAPSHOT", "on")
    return tmp_path


def _word(path, heading_level=2, title="Exported"):
    doc = WordDocument()
    doc.core_properties.title = title
    doc.add_heading("Introduction", level=1)
    doc.add_heading("Background", level=heading_level)
    doc.add_paragraph("The background paragraph explains what came before this study.")
    doc.save(str(path))
    return path


@pytest.fixture
def job(tmp_path):
    """A COMPLETE job whose exported DOCX exists on disk."""
    from src.api.jobs import Job, JobStatus, _jobs, _lock
    from src.models.contracts import ProcessingStatus
    from src.pipeline.phase1_pipeline import PipelineResult

    exported = _word(tmp_path / "exported.docx")
    result = PipelineResult(source_pdf_path="book.pdf", success=True, status=ProcessingStatus.VALIDATED,
                            duration_seconds=0.1, document=None, docx_path=exported)
    job_id = uuid.uuid4().hex
    with _lock:
        _jobs[job_id] = Job(job_id=job_id, filename="book.pdf", pdf_path=tmp_path / "book.pdf",
                            status=JobStatus.COMPLETE, created_at=datetime.now(timezone.utc), result=result)
    yield job_id, exported
    with _lock:
        _jobs.pop(job_id, None)


def _upload(client, job_id, path, name="final.docx"):
    with open(path, "rb") as handle:
        return client.post(f"/api/documents/{job_id}/final-docx", files={"file": (name, handle)})


def test_download_snapshots_the_export(client, roots, job):
    job_id, exported = job
    response = client.get(f"/api/documents/{job_id}/download/docx")
    assert response.status_code == 200
    assert (roots / "learning" / job_id / "exported.docx").read_bytes() == exported.read_bytes()


def test_upload_compares_against_the_downloaded_copy(client, roots, job, tmp_path):
    job_id, _ = job
    client.get(f"/api/documents/{job_id}/download/docx")
    final = _word(tmp_path / "final.docx", heading_level=3, title="Remediated")

    response = _upload(client, job_id, final)

    assert response.status_code == 200, response.text
    review = response.json()
    assert review["baseline"] == "downloaded_copy"
    assert review["diff"]["summary"]["by_category"] == {"HEADING_LEVEL_CHANGED": 1, "METADATA_CHANGED": 1}
    assert client.get(f"/api/documents/{job_id}/final-docx").json()["diff"] == review["diff"]


def test_upload_falls_back_to_the_current_export(client, roots, job, tmp_path):
    job_id, _ = job
    response = _upload(client, job_id, _word(tmp_path / "final.docx"))
    assert response.status_code == 200
    assert response.json()["baseline"] == "current_export"
    assert response.json()["diff"]["edits"] == []


def test_non_docx_uploads_are_refused(client, roots, job, tmp_path):
    job_id, _ = job
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    assert _upload(client, job_id, pdf, name="x.pdf").status_code == 400
    fake = tmp_path / "fake.docx"
    fake.write_bytes(b"not a zip at all")
    response = _upload(client, job_id, fake)
    assert response.status_code == 400
    assert "Word" in response.json()["detail"]


def test_unknown_job_is_404(client, roots, tmp_path):
    assert _upload(client, "nope", _word(tmp_path / "final.docx")).status_code == 404
    assert client.get("/api/documents/nope/final-docx").status_code == 404


def test_delete_forgets_the_review(client, roots, job, tmp_path):
    job_id, _ = job
    _upload(client, job_id, _word(tmp_path / "final.docx", heading_level=3))
    assert client.delete(f"/api/documents/{job_id}/final-docx").status_code == 204
    assert client.get(f"/api/documents/{job_id}/final-docx").status_code == 404
    assert client.delete(f"/api/documents/{job_id}/final-docx").status_code == 404


def test_insights_and_dataset_reflect_uploads(client, roots, job, tmp_path):
    job_id, _ = job
    _upload(client, job_id, _word(tmp_path / "final.docx", heading_level=3))

    insights = client.get("/api/learning/insights").json()
    assert insights["after_export"]["documents"] == 1
    assert insights["after_export"]["categories"][0]["category"] == "HEADING_LEVEL_CHANGED"
    assert insights["priorities"][0]["source"] == "after_export"

    response = client.get("/api/learning/dataset")
    assert response.headers["content-type"].startswith("application/x-ndjson")
    rows = [json.loads(line) for line in response.text.splitlines()]
    assert rows and all(row["kind"] == "post_export_edit" for row in rows)


def test_checklist_delta_reports_fixes_and_regressions():
    from src.api.learning_routes import _checklist_delta

    def report(statuses):
        return {"passed": True, "summary": {}, "results": [
            {"item_id": item, "requirement": f"req {item}", "status": status, "evidence": ""}
            for item, status in statuses.items()]}

    delta = _checklist_delta(report({"A": "fail", "B": "pass", "C": "warn"}),
                             report({"A": "pass", "B": "fail", "C": "warn"}))
    assert [i["item_id"] for i in delta["fixed_by_remediator"]] == ["A"]
    assert [i["item_id"] for i in delta["regressed_by_remediator"]] == ["B"]
