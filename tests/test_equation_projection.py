"""Equations in both projections (docs/EQUATION_DESIGN.md §2, §4, §5).

Markdown spells the LaTeX; the DOCX turns it into native Word equations, with
the number outside the box (Eq 4), functions from the equation tool (Eq 3),
and a fallback that never drops an equation it cannot convert.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from lxml import etree

from src.docx.docx_generator import generate_docx
from src.equations import omml
from src.markdown.markdown_builder import build_markdown
from src.mathpix.ingestor import MathpixImportProvider
from src.models.contracts import Document
from src.models.equation import EquationStatus
from src.models.metadata import Metadata
from src.models.page import Page

_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
needs_pandoc = pytest.mark.skipif(omml.pandoc_path() is None, reason="pandoc not installed")


def _doc(mmd: str, tmp_path: Path, pages: int = 1) -> Document:
    mmd_file = tmp_path / "t.mmd"
    mmd_file.write_text(mmd, encoding="utf-8")
    document = Document(
        source_pdf_path="t.pdf",
        metadata=Metadata(filename="t.pdf", page_count=pages),
        pages=[Page(page_number=i + 1) for i in range(pages)],
    )
    return MathpixImportProvider().import_document(document, mmd_path=mmd_file)


def _docx(document: Document, tmp_path: Path):
    markdown = build_markdown(document)
    path = generate_docx(document, markdown, tmp_path / "out.docx")
    with zipfile.ZipFile(path) as package:
        xml = package.read("word/document.xml")
    return etree.fromstring(xml), markdown


def _paragraph_texts(root):
    return ["".join(p.itertext()) for p in root.iter(f"{{{_W}}}p")]


def _omml_count(root, tag="oMath"):
    return len(list(root.iter(f"{{{_M}}}{tag}")))


class TestMarkdown:
    def test_converted_display_equation_is_spelled_with_its_number(self, tmp_path):
        doc = _doc("\\begin{equation}\nE=m c^{2} \\tag{6.1}\n\\end{equation}\n", tmp_path)
        markdown = build_markdown(doc)
        eq = doc.equations[0]
        assert f"<!-- equation-id: {eq.id} -->" in markdown
        assert "$$E=m c^{2} \\tag{6.1}$$" in markdown

    def test_plain_display_equation_is_a_text_line(self, tmp_path):
        doc = _doc("\\begin{equation*}\nF=m a\n\\end{equation*}\n", tmp_path)
        markdown = build_markdown(doc)
        assert "F = ma" in markdown
        assert "$$" not in markdown

    def test_inline_equation_is_spelled_at_its_anchor(self, tmp_path):
        doc = _doc("Energy is $x^{2}+y^{2}=r^{2}$ here.\n", tmp_path)
        assert "Energy is $x^{2}+y^{2}=r^{2}$ here." in build_markdown(doc)

    def test_plain_inline_equation_stays_text(self, tmp_path):
        doc = _doc("Also \\(v=u+a t\\) inline.\n", tmp_path)
        markdown = build_markdown(doc)
        assert "Also v = u + at inline." in markdown
        assert "$" not in markdown

    def test_inline_equation_and_note_marker_are_both_spelled(self, tmp_path):
        doc = _doc("Claim $x^{2}+y^{2}$ holds ${ }^{1}$.\n\n1. The note body.\n", tmp_path)
        markdown = build_markdown(doc)
        assert "$x^{2}+y^{2}$" in markdown
        assert "[^" in markdown

    def test_edge_text_is_a_paragraph_beside_the_equation(self, tmp_path):
        doc = _doc("$$\n\\text{where } a^{2}=b^{2} \\text{ for all}\n$$\n", tmp_path)
        lines = [line for line in build_markdown(doc).splitlines() if line.strip()]
        equation_line = next(i for i, line in enumerate(lines) if line.startswith("$$"))
        assert lines[equation_line - 1].startswith("<!-- equation-id:")
        assert "where" in lines[equation_line - 2]
        assert "for all" in lines[equation_line + 1]

    def test_a_reviewed_equation_is_spelled_from_the_model(self, tmp_path):
        doc = _doc("$$\nx^{2}\n$$\n", tmp_path)
        doc.equations[0].latex = "y^{3}"
        assert "$$y^{3}$$" in build_markdown(doc)

    def test_a_stale_inline_anchor_leaves_the_text_alone(self, tmp_path):
        doc = _doc("Energy is $x^{2}+y^{2}$ here.\n", tmp_path)
        doc.paragraphs[0].text = "A reviewer rewrote this sentence."
        assert "A reviewer rewrote this sentence." in build_markdown(doc)


class TestContentStream:
    def test_every_display_equation_is_one_stream_node_and_inline_ones_are_not(self, tmp_path):
        from src.models.content_stream import ContentKind
        from src.structure.content_stream import build_content_stream

        doc = _doc("Energy $x^{2}$ here.\n\n$$\nx^{2}+y^{2}\n$$\n$$\n\\frac{a}{b}\n$$\n", tmp_path)
        nodes = [n for n in build_content_stream(doc).nodes if n.kind is ContentKind.EQUATION]
        display = [e.id for e in doc.equations if e.display]
        assert [n.object_id for n in nodes] == display
        assert len(display) == 2


@needs_pandoc
class TestProjectionCorrectness:
    """PI-14 (src/benchmark/projection.py): the DOCX holds every equation the
    model holds, once, and leaks no LaTeX."""

    def _check(self, doc, tmp_path):
        from src.benchmark.projection import check_equation_projection

        path = generate_docx(doc, build_markdown(doc), tmp_path / "pi.docx")
        return check_equation_projection(doc, path, "t")

    def test_a_faithful_projection_has_no_violations(self, tmp_path):
        doc = _doc(
            "Energy $x^{2}+y^{2}$ here.\n\n$$\n\\frac{a}{b}\n$$\n\\begin{equation*}\nF=m a\n\\end{equation*}\n",
            tmp_path,
        )
        report = self._check(doc, tmp_path)
        assert report.ok, report.violations
        assert report.counts["equations"] == 3

    def test_a_document_without_equations_is_trivially_fine(self, tmp_path):
        doc = _doc("Plain prose only.\n", tmp_path)
        assert self._check(doc, tmp_path).ok

    def test_a_lost_equation_is_reported(self, tmp_path):
        doc = _doc("$$\n\\frac{a}{b}\n$$\n", tmp_path)
        path = generate_docx(doc, build_markdown(doc), tmp_path / "lost.docx")
        from src.benchmark.projection import check_equation_projection

        doc.equations.append(doc.equations[0].model_copy(update={"id": "equation-line-99-0", "source_line": 99}))
        report = check_equation_projection(doc, path, "t")
        assert any(v.kind == "lost_object" for v in report.violations)

    def test_leaked_latex_is_reported(self, tmp_path):
        doc = _doc("$$\n\\frac{a}{b}\n$$\n", tmp_path)
        path = generate_docx(doc, build_markdown(doc), tmp_path / "leak.docx")
        package = DocxDocument(str(path))
        package.add_paragraph("leftover \\frac{a}{b} in prose")
        package.save(str(path))
        from src.benchmark.projection import check_equation_projection

        assert any(
            v.kind == "content_invention" for v in check_equation_projection(doc, path, "t").violations
        )


@needs_pandoc
class TestDocx:
    def test_display_equation_is_a_native_word_equation(self, tmp_path):
        doc = _doc("\\begin{equation*}\nE=m c^{2}\n\\end{equation*}\n\\begin{equation*}\n\\frac{a}{b}\n\\end{equation*}\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert _omml_count(root, "f") == 1

    def test_unnumbered_display_equation_is_a_centred_equation_paragraph(self, tmp_path):
        doc = _doc("$$\nx^{2}+y^{2}=r^{2}\n$$\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert _omml_count(root, "oMathPara") == 1

    def test_number_is_outside_the_equation_and_right_aligned(self, tmp_path):
        doc = _doc("\\begin{equation}\nE=m c^{2} \\tag{6.1}\n\\end{equation}\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        boxed = "".join("".join(m.itertext()) for m in root.iter(f"{{{_M}}}oMath"))
        assert "6.1" not in boxed
        paragraph = next(p for p in root.iter(f"{{{_W}}}p") if "(6.1)" in "".join(p.itertext()))
        assert paragraph.find(f".//{{{_W}}}tab[@{{{_W}}}val='right']") is not None
        assert paragraph.find(f".//{{{_W}}}bookmarkStart") is not None

    def test_inline_equation_sits_inside_its_sentence(self, tmp_path):
        doc = _doc("Energy is $x^{2}+y^{2}=r^{2}$ here.\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        paragraph = next(p for p in root.iter(f"{{{_W}}}p") if p.find(f".//{{{_M}}}oMath") is not None)
        text = "".join(t.text or "" for t in paragraph.iter(f"{{{_W}}}t"))
        assert text.startswith("Energy is") and text.endswith("here.")
        assert "$" not in text

    def test_plain_equation_is_ordinary_text_not_an_equation_box(self, tmp_path):
        doc = _doc("\\begin{equation*}\nF=m a\n\\end{equation*}\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert _omml_count(root) == 0
        assert any("F = ma" in text for text in _paragraph_texts(root))

    def test_named_functions_use_the_equation_tool(self, tmp_path):
        doc = _doc("$$\n\\sin x + \\cos y\n$$\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert _omml_count(root, "func") == 2

    def test_an_unconvertible_equation_keeps_its_latex_and_is_flagged(self, tmp_path):
        doc = _doc("$$\n\\weirdmacro{x}+y^{2}\n$$\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert _omml_count(root) == 0
        assert any("weirdmacro" in text for text in _paragraph_texts(root))
        assert doc.equations[0].status is EquationStatus.FLAGGED

    def test_a_conversion_failure_downgrades_a_converted_equation(self, tmp_path, monkeypatch):
        doc = _doc("$$\nx^{2}+y^{2}\n$$\n", tmp_path)
        monkeypatch.setattr(omml, "pandoc_path", lambda: None)
        root, _ = _docx(doc, tmp_path)
        eq = doc.equations[0]
        assert eq.status is EquationStatus.FLAGGED
        assert any("pandoc" in reason for reason in eq.flag_reasons)
        assert any("x^{2}+y^{2}" in text for text in _paragraph_texts(root))

    def test_description_follows_the_equation_as_a_paragraph(self, tmp_path):
        doc = _doc("$$\nx^{2}+y^{2}\n$$\n", tmp_path)
        doc.equations[0].description = "The sum of two squares."
        root, _ = _docx(doc, tmp_path)
        texts = _paragraph_texts(root)
        assert "The sum of two squares." in texts

    def test_edge_text_is_a_normal_paragraph(self, tmp_path):
        doc = _doc("$$\n\\text{where } a^{2}=b^{2}\n$$\n", tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert "where" in _paragraph_texts(root)

    def test_equation_reference_becomes_a_cross_reference_field(self, tmp_path):
        mmd = (
            "\\begin{equation}\nE=m c^{2} \\tag{6.1} \\label{eq:e}\n\\end{equation}\n\n"
            "See \\eqref{eq:e} for the result.\n"
        )
        doc = _doc(mmd, tmp_path)
        root, _ = _docx(doc, tmp_path)
        instructions = "".join(t.text or "" for t in root.iter(f"{{{_W}}}instrText"))
        assert "REF" in instructions
        names = [b.get(f"{{{_W}}}name") for b in root.iter(f"{{{_W}}}bookmarkStart")]
        assert any(name and instructions.split("REF")[1].split()[0] == name for name in names)

    def test_an_equation_line_is_never_stitched_into_prose_across_a_page_break(self, tmp_path):
        mmd = "Text before\n\n$$\nx^{2}+y^{2}\n$$\n\n\\section*{Next}\n"
        doc = _doc(mmd, tmp_path)
        root, _ = _docx(doc, tmp_path)
        assert _omml_count(root) >= 1

    def test_a_document_without_equations_is_unchanged(self, tmp_path):
        doc = _doc("Plain prose only.\n", tmp_path)
        root, markdown = _docx(doc, tmp_path)
        assert "equation-id" not in markdown
        assert _omml_count(root) == 0
        assert "Plain prose only." in _paragraph_texts(root)
