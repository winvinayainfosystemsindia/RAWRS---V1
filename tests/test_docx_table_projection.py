"""P4c-1 — the DOCX projection renders tables from the traversal, not the pipes.

DOCX used to learn *that* a table was there by accumulating ``|...|`` lines and
*which* table it was from a ``<!-- table-id: ... -->`` comment. The grid it then
built already came from the ``Table`` model — so the identity was real, but its
transport was a rendered string, and the pipe-parsing fallback sat one missing
comment away from taking over.

Now a ``TABLE`` node names the table and the model renders it. The decisive
tests hand ``generate_docx`` a **stale** Markdown string, built before the model
was edited: a projection still reading pipes renders the old cells.
"""

from pathlib import Path
from typing import List, Optional

from docx import Document as DocxDocument

from src.docx.docx_generator import generate_docx
from src.markdown.markdown_builder import build_markdown
from src.models.bounding_box import BoundingBox
from src.models.contracts import (
    Document,
    Metadata,
    Page,
    Paragraph,
    Table,
    TableCell,
    TableRow,
    TextBlock,
)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _block(order: int, text: str, page: int = 1) -> TextBlock:
    return TextBlock(
        page_number=page,
        text=text,
        bbox=BoundingBox(x0=72, y0=40 + order * 20, x1=400, y1=56 + order * 20),
        order=order,
    )


def _table(
    table_id: str,
    grid: List[List[str]],
    page: int = 1,
    caption: Optional[str] = None,
    summary: Optional[str] = None,
    header_row: bool = False,
    source_block_ids: Optional[List[str]] = None,
) -> Table:
    rows = [
        TableRow(
            cells=[
                TableCell(
                    text=text,
                    row_index=r,
                    col_index=c,
                    is_header=header_row and r == 0,
                )
                for c, text in enumerate(row)
            ],
            is_header_row=header_row and r == 0,
        )
        for r, row in enumerate(grid)
    ]
    return Table(
        table_id=table_id,
        page_number=page,
        row_count=len(grid),
        col_count=max(len(r) for r in grid),
        rows=rows,
        caption=caption,
        summary=summary,
        source_block_ids=source_block_ids or [],
    )


def _document(
    tables: List[Table],
    blocks: Optional[List[TextBlock]] = None,
    paragraphs: Optional[List[Paragraph]] = None,
    pages: int = 1,
) -> Document:
    blocks = blocks or []
    return Document(
        source_pdf_path="x.pdf",
        metadata=Metadata(filename="x.pdf", page_count=pages),
        pages=[
            Page(
                page_number=n,
                cleaned_text="\n".join(b.text for b in blocks if b.page_number == n),
            )
            for n in range(1, pages + 1)
        ],
        blocks=blocks,
        paragraphs=paragraphs or [],
        tables=tables,
    )


def _render(document: Document, tmp_path: Path, markdown: Optional[str] = None) -> DocxDocument:
    out = tmp_path / "out.docx"
    generate_docx(document, markdown if markdown is not None else build_markdown(document), out)
    return DocxDocument(str(out))


def _grids(docx: DocxDocument) -> List[List[List[str]]]:
    return [[[c.text for c in row.cells] for row in t.rows] for t in docx.tables]


def _body_texts(docx: DocxDocument) -> List[str]:
    return [p.text for p in docx.paragraphs if p.text.strip()]


def _prose_with_a_table(text: str = "Body prose here.") -> tuple:
    """A block-bearing page, so the document is a realistic stream page."""
    block = _block(0, text)
    paragraph = Paragraph(page_number=1, text=text, source_block_ids=[block.block_id])
    return block, paragraph


class TestOneNodeOneTable:
    def test_single_table_renders_once(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["a", "b"], ["c", "d"]])
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        assert _grids(docx) == [[["a", "b"], ["c", "d"]]]

    def test_no_pipe_text_leaks_into_the_body(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["a", "b"], ["c", "d"]])
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        assert not any("|" in text for text in _body_texts(docx))

    def test_table_survives_a_document_with_no_prose(self, tmp_path: Path) -> None:
        table = _table("t1", [["only", "cells"]])
        docx = _render(_document([table]), tmp_path)
        assert _grids(docx) == [[["only", "cells"]]]


class TestOrderAndIdentity:
    def test_multiple_tables_keep_stream_order(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        tables = [
            _table("t1", [["first"]]),
            _table("t2", [["second"]]),
            _table("t3", [["third"]]),
        ]
        docx = _render(_document(tables, [block], [paragraph]), tmp_path)
        assert _grids(docx) == [[["first"]], [["second"]], [["third"]]]

    def test_tables_stay_on_their_own_pages(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Page one prose."), _block(0, "Page two prose.", page=2)]
        paragraphs = [
            Paragraph(page_number=b.page_number, text=b.text, source_block_ids=[b.block_id])
            for b in blocks
        ]
        tables = [_table("t1", [["one"]], page=1), _table("t2", [["two"]], page=2)]
        docx = _render(_document(tables, blocks, paragraphs, pages=2), tmp_path)
        assert _grids(docx) == [[["one"]], [["two"]]]

    def test_identical_grids_are_two_tables(self, tmp_path: Path) -> None:
        """Identity is the table_id, not the text — two tables reading alike
        are two objects and must both render."""
        block, paragraph = _prose_with_a_table()
        tables = [_table("t1", [["same"]]), _table("t2", [["same"]])]
        docx = _render(_document(tables, [block], [paragraph]), tmp_path)
        assert _grids(docx) == [[["same"]], [["same"]]]

    def test_table_appears_after_the_page_prose(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table("The prose above the table.")
        table = _table("t1", [["cell"]])
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        children = [child.tag for child in docx.element.body]
        assert children.index(f"{_W}tbl") > children.index(f"{_W}p")


class TestSemanticsFromTheModel:
    def test_caption_renders_from_the_model(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["x"]], caption="Table 1. A caption.")
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        assert "Table 1. A caption." in _body_texts(docx)

    def test_summary_renders_from_the_model(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["x"]], summary="Two columns of counts.")
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        # _add_table_summary prefixes its own label — projection syntax, not
        # model text, so the assertion is on containment rather than equality.
        assert any("Two columns of counts." in text for text in _body_texts(docx))

    def test_header_row_emits_tbl_header(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["H1", "H2"], ["a", "b"]], header_row=True)
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        rows = docx.tables[0].rows
        assert rows[0]._tr.find(f"{_W}trPr/{_W}tblHeader") is not None
        assert rows[1]._tr.find(f"{_W}trPr/{_W}tblHeader") is None

    def test_header_cells_are_bold(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["H1", "H2"], ["a", "b"]], header_row=True)
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        header = docx.tables[0].rows[0].cells[0].paragraphs[0].runs[0]
        body = docx.tables[0].rows[1].cells[0].paragraphs[0].runs[0]
        assert header.bold is True
        assert body.bold is not True

    def test_merged_cells_emit_a_single_cell(self, tmp_path: Path) -> None:
        """col_span is a merge case the corpus never exercises (0 of 5)."""
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["wide", ""], ["a", "b"]])
        table.rows[0].cells[0].col_span = 2
        docx = _render(_document([table], [block], [paragraph]), tmp_path)
        assert docx.tables[0].cell(0, 0)._tc is docx.tables[0].cell(0, 1)._tc
        assert "wide" in docx.tables[0].cell(0, 0).text


class TestModelBeatsMarkdown:
    def test_stale_markdown_cannot_override_a_cell(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["original"]])
        document = _document([table], [block], [paragraph])
        stale_markdown = build_markdown(document)

        table.rows[0].cells[0].text = "reviewer corrected"

        docx = _render(document, tmp_path, markdown=stale_markdown)
        assert _grids(docx) == [[["reviewer corrected"]]]
        assert "original" in stale_markdown

    def test_reviewer_caption_edit_reaches_docx(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["x"]], caption="Detected caption.")
        document = _document([table], [block], [paragraph])
        stale_markdown = build_markdown(document)

        table.caption = "Reviewer caption."  # the correction rail's effect

        docx = _render(document, tmp_path, markdown=stale_markdown)
        assert "Reviewer caption." in _body_texts(docx)
        assert "Detected caption." not in _body_texts(docx)

    def test_undo_back_to_the_detected_value_also_reaches_docx(self, tmp_path: Path) -> None:
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["detected"]])
        document = _document([table], [block], [paragraph])
        table.rows[0].cells[0].text = "edited"
        edited_markdown = build_markdown(document)

        table.rows[0].cells[0].text = "detected"  # undo restores the model

        docx = _render(document, tmp_path, markdown=edited_markdown)
        assert _grids(docx) == [[["detected"]]]

    def test_markdown_without_the_table_id_comment_still_renders_it(
        self, tmp_path: Path
    ) -> None:
        """The comment was the identity transport. Removing it used to drop the
        table to the pipe-parsing fallback; now the node carries identity."""
        block, paragraph = _prose_with_a_table()
        table = _table("t1", [["a", "b"]], caption="Kept.")
        document = _document([table], [block], [paragraph])
        stripped = "\n".join(
            line
            for line in build_markdown(document).splitlines()
            if not line.strip().startswith("<!-- table-id:")
        )
        docx = _render(document, tmp_path, markdown=stripped)
        assert _grids(docx) == [[["a", "b"]]]
        assert "Kept." in _body_texts(docx)


class TestContainment:
    def test_table_source_blocks_do_not_render_as_prose(self, tmp_path: Path) -> None:
        """A table's own lines are carried by the table, never by a paragraph."""
        cell_blocks = [_block(1, "Alpha"), _block(2, "Beta")]
        prose_block, paragraph = _prose_with_a_table()
        table = _table(
            "t1",
            [["Alpha", "Beta"]],
            source_block_ids=[b.block_id for b in cell_blocks],
        )
        document = _document([table], [prose_block, *cell_blocks], [paragraph])
        docx = _render(document, tmp_path)
        assert _grids(docx) == [[["Alpha", "Beta"]]]
        assert "Alpha" not in _body_texts(docx)
        assert "Beta" not in _body_texts(docx)


class TestFallbackPreserved:
    def test_pipe_markdown_without_a_model_still_renders(self, tmp_path: Path) -> None:
        """Hand-written markdown handed straight to the generator has no
        traversal to read, so the pipe fallback must survive — the only reason
        _add_pipe_table still exists."""
        document = Document(source_pdf_path="x.pdf", metadata=Metadata(filename="x.pdf"))
        markdown = "###### 1\n\n| a | b |\n| --- | --- |\n| c | d |\n"
        docx = _render(document, tmp_path, markdown=markdown)
        assert _grids(docx) == [[["a", "b"], ["c", "d"]]]
