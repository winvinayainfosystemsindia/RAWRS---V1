"""Scanned-page structure: printed page numbers, chapter number, boxes,
heading levels by printed size (src/ocr/scan_structure.py), and the figures,
markdown and DOCX that carry them (image_extractor, markdown_builder,
docx_generator, checklist_audit)."""

from pathlib import Path
from types import SimpleNamespace
from typing import List

import fitz
import numpy as np
import pytest
from docx import Document as WordDocument
from docx.oxml.ns import qn

from src.docx.docx_generator import generate_docx
from src.images.image_extractor import _extract_images_from_pdf
from src.markdown.markdown_builder import BOX_END_MARKER, build_markdown
from src.models.contracts import (
    BoundingBox,
    Callout,
    Document,
    ExtractionMethod,
    Figure,
    Image,
    Metadata,
    OCRLayoutItem,
    Page,
)
from src.ocr import scan_structure
from src.ocr.scan_structure import (
    _apply_printed_labels,
    _detect_boxes,
    _merge_chapter_number,
    apply_scan_structure,
    ascender_height,
    find_panels,
    page_lines,
    split_tiers,
)
from src.validation.checklist_audit import Status, _body_blocks, audit_docx


def _item(label: str, text: str = "", bbox=(72, 72, 300, 90), furniture: bool = False, children=None) -> OCRLayoutItem:
    return OCRLayoutItem(label=label, text=text, bbox=list(bbox), furniture=furniture, children=children or [])


def _page(number: int, items: List[OCRLayoutItem], lines: List[str] = (), levels=None) -> Page:
    return Page(
        page_number=number,
        cleaned_text="\n\n".join(lines),
        extraction_method=ExtractionMethod.DOCLING,
        ocr_layout=items,
        ocr_heading_levels=dict(levels or {}),
    )


def _document(pages: List[Page], pdf: str = "missing.pdf") -> Document:
    return Document(
        source_pdf_path=pdf,
        pages=pages,
        metadata=Metadata(filename=Path(pdf).name, page_count=len(pages), language="en-US", title="T"),
    )


# ------------------------------------------------------------------ page labels


def _folio(text: str) -> OCRLayoutItem:
    return _item("page_footer", text, (300, 740, 320, 755), furniture=True)


class TestPrintedLabels:
    def test_folios_agreeing_on_one_offset_label_every_page(self):
        pages = [_page(n, [_folio(str(27 + n))]) for n in range(1, 5)]
        assert _apply_printed_labels(pages) == 4
        assert [p.printed_label for p in pages] == ["28", "29", "30", "31"]
        assert all(p.label_confidence == 1.0 and not p.label_conflict for p in pages)

    def test_number_inside_a_running_header_counts(self):
        pages = [_page(1, [_item("page_header", "19", furniture=True)]),
                 _page(2, [_item("page_header", "20 Social research strategies", furniture=True)]),
                 _page(3, [_item("page_header", "Social research strategies 21", furniture=True)])]
        _apply_printed_labels(pages)
        assert [p.printed_label for p in pages] == ["19", "20", "21"]

    def test_misread_folio_takes_the_sequence_and_is_flagged(self):
        # O'Leary's page 12: "39" read as "63".
        pages = [_page(n, [_folio("63" if n == 3 else str(27 + n))]) for n in range(1, 5)]
        _apply_printed_labels(pages)
        assert pages[2].printed_label == "30"
        assert pages[2].label_confidence == 0.5 and pages[2].label_conflict

    def test_page_without_folio_is_inferred_not_flagged(self):
        # Bryman's chapter opener prints no number; the book's page is 18.
        pages = [_page(1, [])] + [_page(n, [_folio(str(17 + n))]) for n in range(2, 5)]
        _apply_printed_labels(pages)
        assert pages[0].printed_label == "18"
        assert pages[0].label_confidence == 0.5 and not pages[0].label_conflict

    def test_disagreeing_folios_label_nothing(self):
        pages = [_page(1, [_folio("5")]), _page(2, [_folio("90")]), _page(3, [_folio("12")])]
        assert _apply_printed_labels(pages) == 0
        assert all(p.printed_label is None for p in pages)

    def test_a_label_from_the_text_layer_is_kept(self):
        pages = [_page(n, [_folio(str(10 + n))]) for n in range(1, 4)]
        pages[0].printed_label = "xi"
        _apply_printed_labels(pages)
        assert pages[0].printed_label == "xi"

    def test_running_header_without_a_number_is_not_a_folio(self):
        pages = [_page(n, [_item("page_header", "DEVELOPING YOUR RESEARCH QUESTION", furniture=True)]) for n in (1, 2)]
        assert _apply_printed_labels(pages) == 0


# --------------------------------------------------------------- chapter number


class TestChapterNumber:
    def test_number_text_item_beside_the_title_joins_it(self):
        # O'Leary: a drop numeral taller than the title, beside it.
        page = _page(
            1,
            [_item("text", "3", (129, 75, 200, 165)), _item("section_header", "Developing Your Research Question", (221, 106, 461, 124)),
             _item("text", "Body text follows here.", (72, 200, 500, 260))],
            ["3", "Developing Your Research Question", "Body text follows here."],
            {"Developing Your Research Question": 2},
        )
        assert _merge_chapter_number(page, None) == 1
        assert page_lines(page) == ["3 - Developing Your Research Question", "Body text follows here."]
        assert page.ocr_heading_levels == {"3 - Developing Your Research Question": 2}
        assert page.ocr_layout[1].text == "3 - Developing Your Research Question"

    def test_number_set_as_a_picture_is_read_and_joined(self, monkeypatch):
        # Bryman: a pale "2" Docling kept as a picture with no text.
        monkeypatch.setattr(scan_structure, "_read_picture_text", lambda pdf_page, bbox: "2")
        page = _page(
            1,
            [_item("picture", "", (130, 33, 181, 110)), _item("section_header", "Social research strategies", (120, 139, 376, 230))],
            ["Social research strategies"],
            {"Social research strategies": 2},
        )
        assert _merge_chapter_number(page, {0: object()}) == 1
        assert page_lines(page) == ["2 - Social research strategies"]

    def test_picture_that_is_not_a_number_changes_nothing(self, monkeypatch):
        monkeypatch.setattr(scan_structure, "_read_picture_text", lambda pdf_page, bbox: "Oxford")
        page = _page(1, [_item("picture", "", (130, 33, 181, 110)), _item("section_header", "Title", (120, 139, 376, 160))],
                     ["Title"], {"Title": 2})
        assert _merge_chapter_number(page, {0: object()}) == 0
        assert page_lines(page) == ["Title"]

    def test_figure_sized_picture_is_never_read(self, monkeypatch):
        monkeypatch.setattr(scan_structure, "_read_picture_text", lambda *_: pytest.fail("must not read a figure"))
        page = _page(1, [_item("picture", "", (72, 72, 400, 400)), _item("section_header", "Title", (72, 420, 300, 440))],
                     ["Title"], {"Title": 2})
        assert _merge_chapter_number(page, {0: object()}) == 0

    def test_number_below_the_title_is_not_a_chapter_number(self):
        page = _page(1, [_item("text", "3", (72, 300, 90, 320)), _item("section_header", "Title", (72, 100, 300, 120))],
                     ["3", "Title"], {"Title": 2})
        assert _merge_chapter_number(page, None) == 0

    def test_number_not_followed_by_a_heading_is_left_alone(self):
        page = _page(1, [_item("text", "3", (72, 50, 90, 70)), _item("text", "Prose", (72, 100, 300, 120))], ["3", "Prose"])
        assert _merge_chapter_number(page, None) == 0


# ---------------------------------------------------------------- heading levels


def test_ascender_height_measures_the_band_down_to_the_baseline():
    grey = np.full((60, 40), 255, dtype=np.uint8)
    grey[10:12, 5:8] = 0  # an ascender: little ink
    grey[12:30, 5:35] = 0  # the x-height band: most of the ink
    grey[30:35, 5:8] = 0  # a descender: little ink, not counted
    assert ascender_height(grey, scale=1.0) == 20.0


def test_ascender_height_of_a_blank_crop_is_none():
    assert ascender_height(np.full((10, 10), 255, dtype=np.uint8), scale=1.0) is None


@pytest.mark.parametrize(
    "sizes, expected",
    [
        ([28.8, 15.8, 15.5, 14.0, 13.3, 10.1, 9.9, 9.7], 5),  # Bryman: title, H2 tier, H3 tier
        ([13.1, 10.1, 9.9, 9.7, 9.6, 9.4, 9.0], None),  # O'Leary: one tier, capitals decide
        ([30.0, 12.0, 11.8], None),  # a lone large heading is not a tier
        ([14.0, 13.9, 9.0], None),  # nor is a lone small one
        ([14.0, 13.9, 9.0, 8.9], 2),
    ],
)
def test_split_tiers(sizes, expected):
    assert split_tiers(sizes) == expected


def _headings_pdf(path: Path, headings) -> Path:
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    for text, size, y in headings:
        page.insert_text((72, y), text, fontsize=size, fontname="helv")
    pdf.save(str(path))
    return path


def _heading_items(spec) -> List[OCRLayoutItem]:
    return [_item("section_header", text, (70, y - size, 72 + len(text) * size * 0.6, y + size * 0.3)) for text, size, y in spec]


def test_headings_are_levelled_by_printed_size(tmp_path):
    spec = [("Social research strategies", 32, 80), ("Theory and research", 20, 160), ("Epistemology", 20, 240),
            ("What type of theory", 11, 320), ("Deductive theory", 11, 400)]
    pdf = _headings_pdf(tmp_path / "h.pdf", spec)
    page = _page(1, _heading_items(spec), [s[0] for s in spec], {s[0]: 2 for s in spec})
    counts = apply_scan_structure(_document([page], str(pdf)))
    assert counts["headings_tiered"] == 5
    assert page.ocr_heading_levels == {
        "Social research strategies": 1, "Theory and research": 2, "Epistemology": 2,
        "What type of theory": 3, "Deductive theory": 3,
    }


def test_headings_of_one_size_keep_their_level(tmp_path):
    spec = [("Developing Your Question", 12, 80), ("Chapter Preview", 12, 160), ("Defining", 12, 240), ("Boundaries", 12, 320)]
    pdf = _headings_pdf(tmp_path / "h.pdf", spec)
    page = _page(1, _heading_items(spec), [s[0] for s in spec], {s[0]: 2 for s in spec})
    assert apply_scan_structure(_document([page], str(pdf)))["headings_tiered"] == 0
    assert set(page.ocr_heading_levels.values()) == {2}


# ------------------------------------------------------------------------ boxes


def test_find_panels_takes_a_gradient_panel_to_its_faded_edge():
    grey = np.full((792, 612), 255, dtype=np.uint8)
    grey[300:420, 60:560] = np.linspace(230, 253, 500).astype(np.uint8)  # fades to almost white
    grey[340:350, 100:500] = 0  # a line of text inside
    (x0, y0, x1, y1), = find_panels(grey)
    assert (y0, y1) == (300, 420)
    assert x0 <= 62 and x1 >= 555


def test_find_panels_ignores_a_white_page_with_text():
    grey = np.full((792, 612), 255, dtype=np.uint8)
    grey[100:110, 72:540] = 0
    grey[130:140, 72:540] = 0
    assert find_panels(grey) == []


def _box_pdf(path: Path, rect=(60, 300, 560, 460)) -> Path:
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    page.draw_rect(fitz.Rect(*rect), color=None, fill=(0.85, 0.85, 0.85))
    pdf.save(str(path))
    return path


def _box_page() -> Page:
    items = [
        _item("text", "Intro paragraph.", (72, 100, 540, 200)),
        _item("section_header", "Box 1.1 An example", (80, 310, 400, 330)),
        _item("section_header", "INNER HEADING", (80, 340, 300, 355)),
        _item("text", "Box body text.", (80, 360, 540, 440)),
        _item("text", "After the box.", (72, 500, 540, 600)),
    ]
    return _page(1, items, [i.text for i in items], {"Box 1.1 An example": 2, "INNER HEADING": 2})


def test_a_shaded_panel_becomes_a_box_around_exactly_its_lines(tmp_path):
    pdf = fitz.open(str(_box_pdf(tmp_path / "b.pdf")))
    page = _box_page()
    (box,) = _detect_boxes(page, pdf[0], 0)
    assert (box.page_number, box.first_line, box.last_line) == (1, 1, 3)
    assert box.label == "Box 1.1 An example" and box.callout_type == "box"
    # Only the title stays a heading.
    assert page.ocr_heading_levels == {"Box 1.1 An example": 2}


def test_untitled_panel_is_a_box_labelled_by_its_first_line(tmp_path):
    pdf = fitz.open(str(_box_pdf(tmp_path / "b.pdf", rect=(300, 300, 560, 460))))
    items = [_item("text", "Left column.", (72, 300, 290, 460)), _item("text", "A pull quote.", (310, 320, 550, 440))]
    page = _page(1, items, ["Left column.", "A pull quote."])
    (box,) = _detect_boxes(page, pdf[0], 0)
    assert (box.first_line, box.last_line, box.label) == (1, 1, "A pull quote.")


def test_panel_under_a_picture_is_the_figure_not_a_box(tmp_path):
    pdf = fitz.open(str(_box_pdf(tmp_path / "b.pdf")))
    page = _box_page()
    page.ocr_layout.append(_item("picture", "", (60, 300, 560, 460)))
    assert _detect_boxes(page, pdf[0], 0) == []


def test_box_whose_lines_are_not_contiguous_is_skipped(tmp_path):
    pdf = fitz.open(str(_box_pdf(tmp_path / "b.pdf")))
    page = _box_page()
    page.cleaned_text = "\n\n".join(["Box 1.1 An example", "Intro paragraph.", "INNER HEADING", "Box body text.", "After the box."])
    assert _detect_boxes(page, pdf[0], 0) == []
    assert "INNER HEADING" in page.ocr_heading_levels


def test_born_digital_pages_are_untouched(tmp_path):
    page = Page(page_number=1, cleaned_text="Text", extraction_method=ExtractionMethod.DIRECT_TEXT_EXTRACTION)
    document = _document([page], str(_box_pdf(tmp_path / "b.pdf")))
    assert apply_scan_structure(document) == {"page_labels": 0, "chapter_numbers": 0, "boxes": 0, "headings_tiered": 0}
    assert document.callouts == [] and page.printed_label is None


# ----------------------------------------------------------------- Docling layout


class _Prov:
    def __init__(self, bbox):
        self.page_no = 1
        box = SimpleNamespace(l=bbox[0], t=bbox[1], r=bbox[2], b=bbox[3])
        self.bbox = SimpleNamespace(to_top_left_origin=lambda page_height: box)


def _docling_item(label, text, bbox, layer):
    return SimpleNamespace(label=SimpleNamespace(value=label), text=text, prov=[_Prov(bbox)], content_layer=layer)


def test_layout_items_keep_furniture_and_put_picture_text_under_the_picture():
    from docling_core.types.doc.document import ContentLayer

    from src.ocr.docling_engine import _layout_items

    entries = [
        (_docling_item("page_header", "24 Social research strategies", (51, 16, 192, 33), ContentLayer.FURNITURE), 1),
        (_docling_item("text", "  Body text. ", (70, 60, 300, 100), ContentLayer.BODY), 1),
        (_docling_item("picture", "", (315, 59, 593, 436), ContentLayer.BODY), 1),
        (_docling_item("section_header", "Figure 2.1", (314, 58, 429, 87), ContentLayer.BODY), 2),
        (_docling_item("text", "The process of deduction", (326, 94, 451, 106), ContentLayer.BODY), 2),
        (_docling_item("text", "After the figure.", (320, 585, 536, 724), ContentLayer.BODY), 1),
    ]
    doc = SimpleNamespace(
        iterate_items=lambda **kwargs: iter(entries),
        pages={1: SimpleNamespace(size=SimpleNamespace(height=792))},
    )
    items = _layout_items(doc)
    assert [(i["label"], i["text"], i["furniture"]) for i in items] == [
        ("page_header", "24 Social research strategies", True),
        ("text", "Body text.", False),
        ("picture", "", False),
        ("text", "After the figure.", False),
    ]
    assert items[2]["children"] == ["Figure 2.1", "The process of deduction"]


def test_layout_items_survive_an_unreadable_document():
    from src.ocr.docling_engine import _layout_items

    assert _layout_items(SimpleNamespace()) is None


# ------------------------------------------------------------------------ figures


def _figure_pdf(path: Path) -> Path:
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    page.draw_rect(fitz.Rect(320, 60, 590, 430), color=(0, 0, 0), width=2)
    page.draw_rect(fitz.Rect(70, 500, 300, 600), color=(0, 0, 0), width=2)
    pdf.save(str(path))
    return path


def test_scanned_page_figures_are_cropped_from_its_layout(tmp_path):
    pdf = _figure_pdf(tmp_path / "f.pdf")
    page = _page(1, [
        _item("picture", "", (65, 60, 113, 110)),  # an icon
        _item("picture", "", (315, 59, 593, 436), children=["Figure 2.1", "The process of deduction", "1. Theory"]),
        _item("table", "", (70, 500, 300, 600)),
        _item("picture", "", (66, 490, 310, 610)),  # the table's frame
        _item("caption", "FIGURE 3.2 FROM TOPICS TO QUESTIONS", (70, 640, 400, 652)),
        _item("picture", "", (70, 655, 300, 760)),
    ], ["FIGURE 3.2 FROM TOPICS TO QUESTIONS"])
    images = _extract_images_from_pdf(_document([page], str(pdf)), tmp_path / "out")
    assert len(images) == 2
    first, second = images
    assert Path(first.file_path).is_file() and first.file_path.endswith(".png")
    assert (first.bbox.x0, first.bbox.y1) == (315, 436)
    assert first.figure.caption == "Figure 2.1 The process of deduction"
    assert first.figure.caption_source_text is None  # never in the page text
    assert second.figure.caption == "FIGURE 3.2 FROM TOPICS TO QUESTIONS"
    assert second.figure.caption_source_text == "FIGURE 3.2 FROM TOPICS TO QUESTIONS"


def test_scanned_page_embedded_images_are_not_figures(tmp_path):
    pdf = fitz.open()
    page = pdf.new_page(width=612, height=792)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 200), 0)
    pix.clear_with(128)
    page.insert_image(fitz.Rect(100, 100, 300, 300), pixmap=pix)
    pdf.save(str(tmp_path / "e.pdf"))
    scanned = _page(1, [_item("text", "Only text here.", (72, 400, 500, 420))], ["Only text here."])
    assert _extract_images_from_pdf(_document([scanned], str(tmp_path / "e.pdf")), tmp_path / "out") == []
    native = Page(page_number=1, cleaned_text="x", extraction_method=ExtractionMethod.DIRECT_TEXT_EXTRACTION)
    assert len(_extract_images_from_pdf(_document([native], str(tmp_path / "e.pdf")), tmp_path / "out2")) == 1


# ----------------------------------------------------------------- markdown, DOCX


def _rendered_page() -> Document:
    lines = ["Opening paragraph.", "Box 1.1 An example", "Box body text.", "FIGURE 1 A DIAGRAM", "Closing paragraph."]
    page = _page(1, [
        _item("text", lines[0], (72, 72, 540, 120)),
        _item("section_header", lines[1], (80, 150, 400, 170)),
        _item("text", lines[2], (80, 175, 540, 240)),
        _item("picture", "", (100, 300, 400, 500)),
        _item("caption", lines[3], (100, 505, 400, 520)),
        _item("text", lines[4], (72, 560, 540, 600)),
    ], lines)
    image = Image(
        image_id="fig1", page_number=1, file_path="fig1.png", bbox=BoundingBox(x0=100, y0=300, x1=400, y1=500),
        figure=Figure(caption=lines[3], caption_source_text=lines[3], alt_text="A diagram."),
    )
    document = _document([page])
    document.images = [image]
    document.callouts = [Callout(callout_type="box", label=lines[1], document_order=0, page_number=1, first_line=1, last_line=2)]
    return document


def test_markdown_fences_the_box_and_places_the_figure_where_the_layout_has_it():
    markdown = build_markdown(_rendered_page())
    body = [block for block in markdown.split("\n\n") if block.strip()]
    start = body.index("<!-- box-start: callout-0 -->")
    assert body[start - 1] == "Opening paragraph."
    assert body[start + 1:start + 4] == ["Box 1.1 An example", "Box body text.", BOX_END_MARKER]
    figure = body.index("![A diagram.](fig1.png)")
    assert body[figure + 1] == "*FIGURE 1 A DIAGRAM*"
    assert body[figure + 2] == "Closing paragraph."
    assert markdown.count("FIGURE 1 A DIAGRAM") == 1


def test_docx_sets_a_box_in_a_one_cell_table_with_its_heading_inside(tmp_path):
    markdown = "\n\n".join([
        "###### 1", "Before the box.", "<!-- box-start: callout-0 -->", "### Box 1.1 An example",
        "Box body text.", BOX_END_MARKER, "After the box.",
    ])
    path = generate_docx(_document([]), markdown, output_path=tmp_path / "box.docx")
    word = WordDocument(str(path))
    (table,) = word.tables
    assert (len(table.rows), len(table.columns)) == (1, 1)
    inside = [(p.style.name, p.text) for p in table.rows[0].cells[0].paragraphs]
    # Levels are re-ranked without gaps, so the only content heading is H1.
    assert inside == [("Heading 1", "Box 1.1 An example"), ("Normal", "Box body text.")]
    order = [child.tag.split("}")[1] for child in word.element.body if child.tag != qn("w:sectPr")]
    texts = [p.text for p in word.paragraphs]
    assert order.index("tbl") == texts.index("Before the box.") + 1
    assert texts.index("After the box.") == texts.index("Before the box.") + 1
    # The audit reads inside the box and does not take it for a data table.
    opened = [b.paragraph.text for b in _body_blocks(word, open_boxes=True) if b.kind == "p"]
    assert opened[-3:] == ["Box 1.1 An example", "Box body text.", "After the box."]
    results = {r.item_id: r for r in audit_docx(path).results}
    assert results["SUB-TBL-HDR"].status is Status.NOT_APPLICABLE


def test_box_finds_list_lines_written_with_their_marker(tmp_path):
    # Docling's text export writes "- item"; the layout item's text has no marker.
    pdf = fitz.open(str(_box_pdf(tmp_path / "b.pdf")))
    items = [
        _item("section_header", "Box 3.2 The checklist", (80, 310, 400, 330)),
        _item("list_item", "Will the question hold my interest?", (80, 340, 540, 355)),
        _item("list_item", "Is it doable?", (80, 360, 540, 375)),
    ]
    lines = ["Box 3.2 The checklist", "- Will the question hold my interest?", "- Is it doable?"]
    page = _page(1, items, lines, {"Box 3.2 The checklist": 2})
    (box,) = _detect_boxes(page, pdf[0], 0)
    assert (box.first_line, box.last_line) == (0, 2)


# ------------------------------------------------- review: repeated lines, fences


def test_box_lines_repeated_outside_the_box_resolve_to_their_own_occurrence(tmp_path):
    pdf = fitz.open(str(_box_pdf(tmp_path / "b.pdf")))
    items = [
        _item("section_header", "Summary", (72, 100, 300, 120)),
        _item("text", "Yes", (72, 130, 300, 145)),
        _item("section_header", "Box 1.1 An example", (80, 310, 400, 330)),
        _item("section_header", "Summary", (80, 335, 300, 350)),
        _item("text", "Yes", (80, 360, 540, 380)),
    ]
    lines = ["Summary", "Yes", "Box 1.1 An example", "Summary", "Yes"]
    page = _page(1, items, lines, {"Summary": 2, "Box 1.1 An example": 2})
    (box,) = _detect_boxes(page, pdf[0], 0)
    assert (box.first_line, box.last_line) == (2, 4)
    # "Summary" is still a heading above the box, so it keeps its level.
    assert page.ocr_heading_levels == {"Summary": 2, "Box 1.1 An example": 2}


def test_a_line_already_in_a_box_is_not_boxed_twice():
    from src.ocr.scan_structure import _box_callout

    items = [_item("text", "A", (80, 310, 400, 330)), _item("text", "B", (80, 340, 400, 360))]
    page = _page(1, items, ["A", "B"])
    found = {id(items[0]): 0, id(items[1]): 1}
    assert _box_callout(page, items, found, ["A", "B"], {1}, 0) is None
    assert _box_callout(page, items, found, ["A", "B"], set(), 0).last_line == 1


def test_figure_anchor_skips_an_earlier_copy_of_the_following_line():
    document = _rendered_page()
    page = document.pages[0]
    # The line after the figure also appears before it.
    page.cleaned_text = "\n\n".join(["Closing paragraph.", "Box 1.1 An example", "Box body text.",
                                     "FIGURE 1 A DIAGRAM", "Closing paragraph."])
    page.ocr_layout[0] = _item("text", "Closing paragraph.", (72, 72, 540, 120))
    del page.ocr_layout[4]  # the caption item: the picture's next item is the repeated line
    document.callouts = []
    body = [b for b in build_markdown(document).split("\n\n") if b.strip()]
    figure = body.index("![A diagram.](fig1.png)")
    assert body[figure - 1] == "Box body text." and body[figure + 2] == "Closing paragraph."


def test_caption_line_with_a_list_marker_is_suppressed_once(tmp_path):
    pdf = _figure_pdf(tmp_path / "f.pdf")
    page = _page(1, [
        _item("caption", "FIGURE 3.2 FROM TOPICS", (70, 640, 400, 652)),
        _item("picture", "", (70, 655, 300, 760)),
    ], ["- FIGURE 3.2 FROM TOPICS"])
    (image,) = _extract_images_from_pdf(_document([page], str(pdf)), tmp_path / "out")
    assert image.figure.caption == "FIGURE 3.2 FROM TOPICS"
    assert image.figure.caption_source_text == "- FIGURE 3.2 FROM TOPICS"


def test_an_unclosed_box_fence_does_not_swallow_the_next_page(tmp_path):
    from src.markdown.markdown_builder import PAGE_BREAK_MARKER

    markdown = "\n\n".join([
        "###### 1", "<!-- box-start: callout-0 -->", "Inside an unclosed box.", PAGE_BREAK_MARKER,
        "###### 2", "Next page.", "<!-- box-start: callout-1 -->", "Boxed.", BOX_END_MARKER,
    ])
    path = generate_docx(_document([]), markdown, output_path=tmp_path / "fence.docx")
    word = WordDocument(str(path))
    (table,) = word.tables
    assert [p.text for p in table.rows[0].cells[0].paragraphs] == ["Boxed."]
    assert "Next page." in [p.text for p in word.paragraphs]


def test_box_table_is_titled_and_a_plain_one_cell_table_stays_a_data_table(tmp_path):
    markdown = "\n\n".join(["###### 1", "<!-- box-start: callout-0 -->", "### Box 2.1 Theory", "Text.", BOX_END_MARKER])
    word = WordDocument(str(generate_docx(_document([]), markdown, output_path=tmp_path / "t.docx")))
    caption = word.tables[0]._tbl.tblPr.find(qn("w:tblCaption"))
    assert caption.get(qn("w:val")) == "Box: Box 2.1 Theory"
    # An unmarked one-cell table is a data table to the audit.
    plain = WordDocument()
    plain.add_table(rows=1, cols=1).rows[0].cells[0].text = "42"
    assert [b.kind for b in _body_blocks(plain, open_boxes=True)] == ["table"]
    # A table inside a box is still checked as a table.
    caption_xml = word.tables[0]._tbl
    caption_xml.findall(qn("w:tr"))[0].findall(qn("w:tc"))[0].append(plain.tables[0]._tbl)
    assert "table" in [b.kind for b in _body_blocks(word, open_boxes=True)]
