"""DR-EQ / DR-EQ-STRUCT: the equation rules of the checklist, audited
(docs/EQUATION_DESIGN.md §1, §5)."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from lxml import etree

from src.docx.docx_generator import generate_docx
from src.equations import omml
from src.markdown.markdown_builder import build_markdown
from src.mathpix.ingestor import MathpixImportProvider
from src.models.contracts import Document
from src.models.metadata import Metadata
from src.models.page import Page
from src.validation.checklist_audit import Status, audit_docx

_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
needs_pandoc = pytest.mark.skipif(omml.pandoc_path() is None, reason="pandoc not installed")


def _run(text: str, nor: bool = False) -> str:
    properties = "<m:rPr><m:nor/></m:rPr>" if nor else ""
    return f"<m:r>{properties}<m:t>{text}</m:t></m:r>"


def _square(base: str, power: str) -> str:
    return f"<m:sSup><m:e>{_run(base)}</m:e><m:sup>{_run(power)}</m:sup></m:sSup>"


def _func(name: str, argument: str) -> str:
    return f"<m:func><m:fName>{_run(name)}</m:fName><m:e>{_run(argument)}</m:e></m:func>"


def _audit(tmp_path: Path, *boxes: str, equations=None):
    document = DocxDocument()
    for box in boxes:
        paragraph = document.add_paragraph()
        paragraph._p.append(etree.fromstring(f'<m:oMath xmlns:m="{_M}">{box}</m:oMath>'))
    path = tmp_path / "audit.docx"
    document.save(str(path))
    report = audit_docx(path, equations=equations)
    return {result.item_id: result for result in report.results}


def _flagged(identifier: str = "equation-line-3-0"):
    return SimpleNamespace(
        id=identifier, status=SimpleNamespace(value="flagged"), flag_reasons=["unsupported command \\x"]
    )


def _converted():
    return SimpleNamespace(id="equation-line-1-0", status=SimpleNamespace(value="converted"), flag_reasons=[])


class TestStructure:
    def test_a_well_formed_equation_passes(self, tmp_path):
        results = _audit(tmp_path, _square("x", "2") + _run("+") + _func("sin", "y"))
        assert results["DR-EQ-STRUCT"].status is Status.PASS

    def test_a_typed_function_name_fails_eq3(self, tmp_path):
        results = _audit(tmp_path, _square("x", "2") + _run("sin") + _run("y"))
        assert results["DR-EQ-STRUCT"].status is Status.FAIL
        assert "Eq 3" in results["DR-EQ-STRUCT"].evidence

    def test_a_function_inside_a_script_base_is_from_the_equation_tool(self, tmp_path):
        inverse = (
            "<m:func><m:fName><m:sSup><m:e>" + _run("sin") + "</m:e><m:sup>" + _run("-1") + "</m:sup></m:sSup>"
            "</m:fName><m:e>" + _run("x") + "</m:e></m:func>"
        )
        assert _audit(tmp_path, inverse)["DR-EQ-STRUCT"].status is Status.PASS

    def test_the_number_inside_the_box_fails_eq4(self, tmp_path):
        results = _audit(tmp_path, _square("x", "2") + _run("(6.1)"))
        assert results["DR-EQ-STRUCT"].status is Status.FAIL
        assert "Eq 4" in results["DR-EQ-STRUCT"].evidence

    def test_several_equations_in_one_box_fail_eq2(self, tmp_path):
        box = "<m:eqArr><m:e>" + _square("x", "2") + "</m:e><m:e>" + _square("y", "2") + "</m:e></m:eqArr>"
        results = _audit(tmp_path, box)
        assert results["DR-EQ-STRUCT"].status is Status.FAIL
        assert "Eq 2" in results["DR-EQ-STRUCT"].evidence

    def test_connecting_text_at_the_edge_fails_eq5(self, tmp_path):
        results = _audit(tmp_path, _run("where", nor=True) + _square("a", "2"))
        assert results["DR-EQ-STRUCT"].status is Status.FAIL
        assert "Eq 5" in results["DR-EQ-STRUCT"].evidence

    def test_a_stray_space_fails_eq1(self, tmp_path):
        results = _audit(tmp_path, _run(" ") + _square("a", "2"))
        assert results["DR-EQ-STRUCT"].status is Status.FAIL
        assert "Eq 1" in results["DR-EQ-STRUCT"].evidence

    def test_a_simple_equation_inside_a_box_fails_eq7(self, tmp_path):
        results = _audit(tmp_path, _run("x") + _run("=") + _run("1"))
        assert results["DR-EQ-STRUCT"].status is Status.FAIL
        assert "Eq 7" in results["DR-EQ-STRUCT"].evidence

    def test_a_document_without_equations_has_nothing_to_audit(self, tmp_path):
        document = DocxDocument()
        document.add_paragraph("Plain prose.")
        path = tmp_path / "plain.docx"
        document.save(str(path))
        results = {r.item_id: r for r in audit_docx(path).results}
        assert results["DR-EQ"].status is Status.NOT_APPLICABLE
        assert "DR-EQ-STRUCT" not in results


class TestSummary:
    def test_without_the_model_a_person_must_compare_with_the_pdf(self, tmp_path):
        results = _audit(tmp_path, _square("x", "2"))
        assert results["DR-EQ"].status is Status.MANUAL

    def test_converted_equations_pass_because_the_round_trip_proved_them(self, tmp_path):
        results = _audit(tmp_path, _square("x", "2"), equations=[_converted()])
        assert results["DR-EQ"].status is Status.PASS
        assert "1 converted" in results["DR-EQ"].evidence

    def test_a_flagged_equation_needs_a_person_and_is_named(self, tmp_path):
        results = _audit(tmp_path, _square("x", "2"), equations=[_converted(), _flagged()])
        assert results["DR-EQ"].status is Status.MANUAL
        assert "equation-line-3-0" in results["DR-EQ"].evidence
        assert results["DR-EQ"].count == 1

    def test_equations_with_no_word_equation_still_report_when_the_model_has_them(self, tmp_path):
        document = DocxDocument()
        document.add_paragraph("No boxes.")
        path = tmp_path / "none.docx"
        document.save(str(path))
        results = {r.item_id: r for r in audit_docx(path, equations=[_flagged()]).results}
        assert results["DR-EQ"].status is Status.MANUAL


@needs_pandoc
class TestOnGeneratedDocuments:
    def _generated(self, mmd: str, tmp_path: Path):
        mmd_file = tmp_path / "t.mmd"
        mmd_file.write_text(mmd, encoding="utf-8")
        document = Document(
            source_pdf_path="t.pdf",
            metadata=Metadata(filename="t.pdf", page_count=1),
            pages=[Page(page_number=1)],
        )
        document = MathpixImportProvider().import_document(document, mmd_path=mmd_file)
        path = generate_docx(document, build_markdown(document), tmp_path / "g.docx")
        return document, {r.item_id: r for r in audit_docx(path, equations=document.equations).results}

    def test_what_rawrs_generates_satisfies_the_structure_rules(self, tmp_path):
        mmd = (
            "Energy $x^{2}+y^{2}$ here.\n\n"
            "\\begin{equation}\n\\sin^{-1} x + \\frac{a}{b} \\tag{6.1}\n\\end{equation}\n\n"
            "\\begin{align*}\nx^{2} & =1 \\\\\ny^{3} & =2\n\\end{align*}\n\n"
            "$$\n\\text{where } a^{2}=b^{2} \\text{ for all}\n$$\n"
            "\\begin{equation*}\nF=m a\n\\end{equation*}\n"
        )
        _, results = self._generated(mmd, tmp_path)
        assert results["DR-EQ-STRUCT"].status is Status.PASS, results["DR-EQ-STRUCT"].evidence
        assert results["DR-EQ"].status is Status.PASS, results["DR-EQ"].evidence

    def test_an_unconvertible_equation_makes_dr_eq_manual(self, tmp_path):
        document, results = self._generated("$$\n\\weirdmacro{x}+y^{2}\n$$\n", tmp_path)
        assert document.equations[0].status.value == "flagged"
        assert results["DR-EQ"].status is Status.MANUAL
        assert document.equations[0].id in results["DR-EQ"].evidence
