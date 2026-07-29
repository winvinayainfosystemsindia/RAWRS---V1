"""F0 — Benchmark Differ tests.

Measurement infrastructure only: no production module is touched. Exercises
the DOCX profiler (against a synthetic .docx), the RAWRS Document profiler
(against a lightweight mock), and the diff/metrics/report layer (against
hand-built profiles).
"""

from types import SimpleNamespace

from src.benchmark.differ import build_report, diff_profiles, metrics
from src.benchmark.profile import (
    BenchmarkProfile,
    Figure,
    Heading,
    Note,
    Table,
    profile_from_docx,
    profile_from_document,
)


def _prof(**over) -> BenchmarkProfile:
    base = dict(source="t", headings=[], paragraph_count=0, tables=[], figures=[],
                footnotes=[], endnotes=[], page_breaks=0, printed_page_labels=[],
                running_header_candidates=[], metadata={"title": None, "language": None})
    base.update(over)
    return BenchmarkProfile(**base)


class TestDocxProfile:
    def test_profile_from_docx_extracts_structure(self, tmp_path):
        from docx import Document as DocxDocument

        d = DocxDocument()
        d.core_properties.title = "My Title"
        d.add_heading("My Title", level=1)
        d.add_heading("Introduction", level=2)
        d.add_paragraph("Body paragraph one.")
        d.add_paragraph("Body paragraph two.")
        d.add_table(rows=2, cols=3)
        path = tmp_path / "sample.docx"
        d.save(str(path))

        prof = profile_from_docx(path)

        assert [h.level for h in prof.headings] == [1, 2]
        assert prof.headings[1].text == "Introduction"
        assert prof.paragraph_count == 2
        assert prof.tables == [Table(rows=2, cols=3)]
        assert prof.metadata["title"] == "My Title"


class TestDocumentProfile:
    def test_profile_from_document_mock(self):
        doc = SimpleNamespace(
            headings=[
                SimpleNamespace(level=SimpleNamespace(value="h2"), text="Intro", is_page_marker=False),
                SimpleNamespace(level=6, text="3", is_page_marker=True),  # page marker excluded
            ],
            blocks=[SimpleNamespace(text="para one"), SimpleNamespace(text="  ")],
            tables=[SimpleNamespace(rows=[SimpleNamespace(cells=[1, 2, 3])])],
            images=[SimpleNamespace(figure=SimpleNamespace(alt_text="a chart"))],
            footnotes=[
                SimpleNamespace(number=1, text="fn body", note_type=SimpleNamespace(value="footnote")),
                SimpleNamespace(number=2, text="en body", note_type=SimpleNamespace(value="endnote")),
            ],
            pages=[SimpleNamespace(printed_label="i"), SimpleNamespace(printed_label=None)],
            metadata=SimpleNamespace(title="Doc Title", language="en-US"),
        )
        prof = profile_from_document(doc)

        assert prof.headings == [Heading(level=2, text="Intro")]  # marker dropped
        assert prof.paragraph_count == 1  # blank block ignored
        assert prof.tables == [Table(rows=1, cols=3)]
        assert prof.figures == [Figure(alt="a chart")]
        assert [n.text for n in prof.footnotes] == ["fn body"]
        assert [n.text for n in prof.endnotes] == ["en body"]
        assert prof.page_breaks == 1
        assert prof.printed_page_labels == ["i"]
        assert prof.metadata == {"title": "Doc Title", "language": "en-US"}


class TestDiffAndMetrics:
    def test_perfect_match_scores_one(self):
        b = _prof(headings=[Heading(1, "Title"), Heading(2, "Intro")])
        m = metrics(b, b)
        assert m["heading"] == {"precision": 1.0, "recall": 1.0, "f1": 1.0}
        assert m["front_page_similarity"] == 1.0
        assert m["running_header_suppression_accuracy"] == 1.0

    def test_missing_and_extra_headings(self):
        b = _prof(headings=[Heading(1, "Title"), Heading(2, "Intro")])
        r = _prof(headings=[Heading(1, "Title"), Heading(2, "Spurious")])
        d = diff_profiles(b, r)
        statuses = {i["status"] for i in d["headings"]["items"]}
        assert "missing" in statuses and "extra" in statuses
        m = metrics(b, r)
        assert m["heading"]["recall"] == 0.5   # 1 of 2 benchmark headings found
        assert m["heading"]["precision"] == 0.5  # 1 of 2 rawrs headings correct

    def test_footnote_placement_penalized_when_misclassified(self):
        b = _prof(footnotes=[Note("1", "a note")])
        good = _prof(footnotes=[Note("1", "a note")])
        bad = _prof(endnotes=[Note("1", "a note")])  # RAWRS put it at the end
        assert metrics(b, good)["footnote_placement_accuracy"] == 1.0
        assert metrics(b, bad)["footnote_placement_accuracy"] == 0.0

    def test_running_header_leak_penalized(self):
        b = _prof(running_header_candidates=[])            # human removed the masthead
        r = _prof(running_header_candidates=["journal of x"])  # RAWRS left it in the body
        assert metrics(b, r)["running_header_suppression_accuracy"] == 0.0

    def test_page_break_fidelity(self):
        assert metrics(_prof(page_breaks=10), _prof(page_breaks=10))["page_break_fidelity"] == 1.0
        assert metrics(_prof(page_breaks=10), _prof(page_breaks=8))["page_break_fidelity"] == 0.8

    def test_build_report_has_corpus_aggregate(self):
        b = _prof(headings=[Heading(1, "T")])
        report = build_report([("doc1", b, b)])
        assert report["schema"] == "rawrs.benchmark-gap-report/v1"
        assert report["corpus"]["document_count"] == 1
        assert report["corpus"]["heading_f1"] == 1.0
        assert report["documents"][0]["document"] == "doc1"
