"""P4b — the DOCX projection renders prose, headings and pages from the stream.

Until this milestone DOCX recovered its structure by parsing the Markdown it
was handed: ``^#{1,6}`` for a heading and its level, ``<!-- pagebreak -->``
for a page boundary, ``**``/``*`` for emphasis, and every remaining line as a
paragraph. Markdown was a semantic intermediate representation, which is the
defect ADR-019 named.

The sharpest tests here hand ``generate_docx`` a **stale** Markdown string —
one built before the model was edited. A projection that still reads Markdown
for meaning renders the old text; one that reads the ContentStream renders the
new. Text matching cannot fake that, and neither can a lucky ordering.
"""

from pathlib import Path
from typing import List, Optional

import pytest
from docx import Document as DocxDocument

from src.docx.docx_generator import generate_docx
from src.markdown.markdown_builder import build_markdown
from src.models.bounding_box import BoundingBox
from src.models.contracts import (
    Document,
    Heading,
    HeadingLevel,
    Metadata,
    Page,
    Paragraph,
    TextBlock,
)
from src.models.inline_format import BOLD_FLAG, ITALIC_FLAG
from src.models.span import Span


def _span(text: str, flags: int = 0) -> Span:
    return Span(
        text=text,
        font_name="Times",
        font_size=12.0,
        font_flags=flags,
        baseline_y=50.0,
        bbox=BoundingBox(x0=72, y0=40, x1=400, y1=56),
    )


def _block(order: int, text: str, page: int = 1, flags: Optional[int] = None) -> TextBlock:
    return TextBlock(
        page_number=page,
        text=text,
        bbox=BoundingBox(x0=72, y0=40 + order * 20, x1=400, y1=56 + order * 20),
        order=order,
        spans=[_span(text, flags)] if flags is not None else [],
    )


def _document(
    blocks: List[TextBlock],
    paragraphs: List[Paragraph],
    headings: Optional[List[Heading]] = None,
    pages: int = 1,
) -> Document:
    return Document(
        source_pdf_path="x.pdf",
        metadata=Metadata(filename="x.pdf", page_count=pages),
        pages=[
            Page(
                page_number=n,
                cleaned_text="\n".join(b.text for b in blocks if b.page_number == n),
            )
            for n in range(1, pages + 1)
        ],
        blocks=blocks,
        paragraphs=paragraphs,
        headings=headings or [],
    )


def _render(document: Document, tmp_path: Path, markdown: Optional[str] = None) -> DocxDocument:
    """Generate and reopen. ``markdown`` defaults to this document's own."""
    out = tmp_path / "out.docx"
    generate_docx(document, markdown if markdown is not None else build_markdown(document), out)
    return DocxDocument(str(out))


def _body(docx: DocxDocument) -> List[str]:
    """Non-heading paragraph text, in document order."""
    return [
        p.text
        for p in docx.paragraphs
        if p.text.strip() and not (p.style.name or "").startswith("Heading")
    ]


def _headings(docx: DocxDocument) -> List[tuple]:
    return [
        ((p.style.name or "").replace("Heading ", ""), p.text)
        for p in docx.paragraphs
        if p.text.strip() and (p.style.name or "").startswith("Heading")
    ]


def _page_breaks(docx: DocxDocument) -> int:
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return sum(
        1 for br in docx.element.body.iter(f"{w}br") if br.get(f"{w}type") == "page"
    )


@pytest.fixture
def simple() -> Document:
    blocks = [_block(0, "First line."), _block(1, "second line."), _block(2, "A closing one.")]
    paragraphs = [
        Paragraph(
            page_number=1,
            text="First line. second line.",
            source_block_ids=[blocks[0].block_id, blocks[1].block_id],
        ),
        Paragraph(page_number=1, text="A closing one.", source_block_ids=[blocks[2].block_id]),
    ]
    return _document(blocks, paragraphs)


class TestParagraphs:
    def test_one_paragraph(self, tmp_path: Path) -> None:
        block = _block(0, "Only prose here.")
        document = _document(
            [block],
            [Paragraph(page_number=1, text="Only prose here.", source_block_ids=[block.block_id])],
        )
        assert _body(_render(document, tmp_path)) == ["Only prose here."]

    def test_multiple_paragraphs_keep_stream_order(
        self, simple: Document, tmp_path: Path
    ) -> None:
        assert _body(_render(simple, tmp_path)) == [
            "First line. second line.",
            "A closing one.",
        ]

    def test_duplicate_paragraph_text_renders_twice(self, tmp_path: Path) -> None:
        """Two paragraphs reading identically are two objects, not one."""
        blocks = [_block(0, "Same words."), _block(1, "Between."), _block(2, "Same words.")]
        paragraphs = [
            Paragraph(page_number=1, text=b.text, source_block_ids=[b.block_id])
            for b in blocks
        ]
        assert _body(_render(_document(blocks, paragraphs), tmp_path)) == [
            "Same words.",
            "Between.",
            "Same words.",
        ]

    def test_corrected_paragraph_text_wins_over_stale_markdown(self, tmp_path: Path) -> None:
        """The §10 end-to-end case: edit -> Document -> stream -> DOCX.

        The Markdown handed to the generator is the one built *before* the
        edit. Only a projection reading the semantic object renders the new
        text; one parsing Markdown renders the old.
        """
        block = _block(0, "The original sentence.")
        paragraph = Paragraph(
            page_number=1, text="The original sentence.", source_block_ids=[block.block_id]
        )
        document = _document([block], [paragraph])
        stale_markdown = build_markdown(document)

        paragraph.text = "The reviewer's corrected sentence."

        body = _body(_render(document, tmp_path, markdown=stale_markdown))
        assert body == ["The reviewer's corrected sentence."]
        assert "The original sentence." in stale_markdown

    def test_bold_paragraph(self, tmp_path: Path) -> None:
        runs = self._runs_for(tmp_path, BOLD_FLAG)
        assert [(r.bold, r.italic) for r in runs] == [(True, None)]

    def test_italic_paragraph(self, tmp_path: Path) -> None:
        runs = self._runs_for(tmp_path, ITALIC_FLAG)
        assert [(r.bold, r.italic) for r in runs] == [(False, True)]

    def test_bold_italic_paragraph(self, tmp_path: Path) -> None:
        runs = self._runs_for(tmp_path, BOLD_FLAG | ITALIC_FLAG)
        assert [(r.bold, r.italic) for r in runs] == [(True, True)]

    def test_plain_paragraph_has_no_emphasis(self, tmp_path: Path) -> None:
        runs = self._runs_for(tmp_path, 0)
        assert [(r.bold, r.italic) for r in runs] == [(False, None)]

    def test_literal_asterisks_are_not_markup(self, tmp_path: Path) -> None:
        """A paragraph ending in ``*****`` keeps all five.

        The old path handed DOCX rendered Markdown and re-parsed ``**`` out
        of it, so a document's own asterisks were consumed as emphasis. One
        corpus paragraph (Aims of Education) lost four characters this way.
        """
        block = _block(0, "Ornamented. *****")
        document = _document(
            [block],
            [Paragraph(page_number=1, text="Ornamented. *****", source_block_ids=[block.block_id])],
        )
        assert _body(_render(document, tmp_path)) == ["Ornamented. *****"]

    @staticmethod
    def _runs_for(tmp_path: Path, flags: int) -> list:
        block = _block(0, "Emphasised prose.", flags=flags)
        document = _document(
            [block],
            [
                Paragraph(
                    page_number=1,
                    text="Emphasised prose.",
                    source_block_ids=[block.block_id],
                )
            ],
        )
        docx = _render(document, tmp_path)
        paragraph = next(
            p
            for p in docx.paragraphs
            if p.text.strip() and not (p.style.name or "").startswith("Heading")
        )
        return [r.font for r in paragraph.runs if r.text]

    def test_paragraph_at_page_boundary_stays_on_its_page(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Page one prose."), _block(0, "Page two prose.", page=2)]
        paragraphs = [
            Paragraph(page_number=b.page_number, text=b.text, source_block_ids=[b.block_id])
            for b in blocks
        ]
        docx = _render(_document(blocks, paragraphs, pages=2), tmp_path)
        assert _body(docx) == ["Page one prose.", "Page two prose."]
        assert _page_breaks(docx) == 1


class TestHeadings:
    @pytest.mark.parametrize("level", [1, 2, 3, 4, 5])
    def test_content_heading_levels_come_from_the_model(
        self, level: int, tmp_path: Path
    ) -> None:
        blocks = [_block(0, "A Heading"), _block(1, "Body prose here.")]
        heading = Heading(
            level=HeadingLevel(level),
            text="A Heading",
            page_number=1,
            document_order=0,
            source_block_id=blocks[0].block_id,
        )
        paragraph = Paragraph(
            page_number=1, text="Body prose here.", source_block_ids=[blocks[1].block_id]
        )
        docx = _render(_document(blocks, [paragraph], [heading]), tmp_path)
        assert (str(level), "A Heading") in _headings(docx)

    def test_h6_page_marker_is_a_heading_6(self, simple: Document, tmp_path: Path) -> None:
        levels = [level for level, _ in _headings(_render(simple, tmp_path))]
        assert "6" in levels

    def test_duplicate_heading_text_renders_twice(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Notes"), _block(1, "Prose."), _block(2, "Notes")]
        headings = [
            Heading(
                level=HeadingLevel.H2,
                text="Notes",
                page_number=1,
                document_order=n,
                source_block_id=blocks[i].block_id,
            )
            for n, i in enumerate((0, 2))
        ]
        paragraph = Paragraph(
            page_number=1, text="Prose.", source_block_ids=[blocks[1].block_id]
        )
        docx = _render(_document(blocks, [paragraph], headings), tmp_path)
        assert [t for _, t in _headings(docx)].count("Notes") == 2

    def test_corrected_heading_wins_over_stale_markdown(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Old Title"), _block(1, "Body prose here.")]
        heading = Heading(
            level=HeadingLevel.H3,
            text="Old Title",
            page_number=1,
            document_order=0,
            source_block_id=blocks[0].block_id,
        )
        paragraph = Paragraph(
            page_number=1, text="Body prose here.", source_block_ids=[blocks[1].block_id]
        )
        document = _document(blocks, [paragraph], [heading])
        stale_markdown = build_markdown(document)

        heading.text = "New Title"
        heading.level = HeadingLevel.H1

        docx = _render(document, tmp_path, markdown=stale_markdown)
        assert ("1", "New Title") in _headings(docx)
        assert "New Title" not in stale_markdown

    def test_heading_adjacent_to_paragraph_keeps_order(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Before."), _block(1, "The Heading"), _block(2, "After.")]
        heading = Heading(
            level=HeadingLevel.H2,
            text="The Heading",
            page_number=1,
            document_order=0,
            source_block_id=blocks[1].block_id,
        )
        paragraphs = [
            Paragraph(page_number=1, text="Before.", source_block_ids=[blocks[0].block_id]),
            Paragraph(page_number=1, text="After.", source_block_ids=[blocks[2].block_id]),
        ]
        docx = _render(_document(blocks, paragraphs, [heading]), tmp_path)
        ordered = [
            p.text
            for p in docx.paragraphs
            if p.text.strip() and p.text != "1"  # drop the page marker
        ]
        assert ordered == ["Before.", "The Heading", "After."]


class TestPageMarkers:
    def test_consecutive_pages_break_between_not_after(self, tmp_path: Path) -> None:
        blocks = [_block(0, f"Page {n} prose.", page=n) for n in (1, 2, 3)]
        paragraphs = [
            Paragraph(page_number=b.page_number, text=b.text, source_block_ids=[b.block_id])
            for b in blocks
        ]
        docx = _render(_document(blocks, paragraphs, pages=3), tmp_path)
        assert _page_breaks(docx) == 2  # three pages, no trailing break

    def test_marker_precedes_its_page_content(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Alpha.", page=1), _block(0, "Beta.", page=2)]
        paragraphs = [
            Paragraph(page_number=b.page_number, text=b.text, source_block_ids=[b.block_id])
            for b in blocks
        ]
        docx = _render(_document(blocks, paragraphs, pages=2), tmp_path)
        ordered = [p.text for p in docx.paragraphs if p.text.strip()]
        assert ordered.index("1") < ordered.index("Alpha.") < ordered.index("2")
        assert ordered.index("2") < ordered.index("Beta.")

    def test_marker_text_comes_from_the_model_not_the_markdown(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Prose.")]
        marker = Heading(
            level=HeadingLevel.H6,
            text="1",
            page_number=1,
            document_order=0,
            is_page_marker=True,
        )
        paragraph = Paragraph(
            page_number=1, text="Prose.", source_block_ids=[blocks[0].block_id]
        )
        document = _document(blocks, [paragraph], [marker])
        stale_markdown = build_markdown(document)

        marker.text = "xvii"  # a reviewed printed label

        docx = _render(document, tmp_path, markdown=stale_markdown)
        assert ("6", "xvii") in _headings(docx)

    def test_marker_around_heading_and_paragraph(self, tmp_path: Path) -> None:
        blocks = [_block(0, "Chapter One"), _block(1, "Its prose.")]
        heading = Heading(
            level=HeadingLevel.H1,
            text="Chapter One",
            page_number=1,
            document_order=0,
            source_block_id=blocks[0].block_id,
        )
        paragraph = Paragraph(
            page_number=1, text="Its prose.", source_block_ids=[blocks[1].block_id]
        )
        docx = _render(_document(blocks, [paragraph], [heading]), tmp_path)
        assert _headings(docx) == [("6", "1"), ("1", "Chapter One")]
        assert _body(docx) == ["Its prose."]


class TestBlocklessSafety:
    """§7 — a page with no TextBlock keeps the Markdown-derived path.

    41 of the corpus' 161 pages have no blocks, so the traversal can place no
    paragraph on them. P4b must leave their text alone rather than render an
    empty page from an empty stream.
    """

    def test_page_without_blocks_still_renders_its_text(self, tmp_path: Path) -> None:
        document = Document(
            source_pdf_path="x.pdf",
            metadata=Metadata(filename="x.pdf", page_count=1),
            pages=[Page(page_number=1, cleaned_text="Recovered by OCR.\nA second line.")],
            blocks=[],
            paragraphs=[],
        )
        body = _body(_render(document, tmp_path))
        assert "Recovered by OCR." in body
        assert "A second line." in body

    def test_mixed_document_keeps_the_blockless_page(self, tmp_path: Path) -> None:
        block = _block(0, "Born-digital prose.", page=1)
        document = Document(
            source_pdf_path="x.pdf",
            metadata=Metadata(filename="x.pdf", page_count=2),
            pages=[
                Page(page_number=1, cleaned_text="Born-digital prose."),
                Page(page_number=2, cleaned_text="Scanned prose."),
            ],
            blocks=[block],
            paragraphs=[
                Paragraph(
                    page_number=1,
                    text="Born-digital prose.",
                    source_block_ids=[block.block_id],
                )
            ],
        )
        body = _body(_render(document, tmp_path))
        assert "Born-digital prose." in body
        assert "Scanned prose." in body

    def test_document_without_paragraphs_is_unchanged(self, tmp_path: Path) -> None:
        """A Document that never ran Stage 5c has no traversal to read."""
        block = _block(0, "Prose with no Paragraph object.")
        document = _document([block], [])
        assert _body(_render(document, tmp_path)) == ["Prose with no Paragraph object."]
