"""Structural profile extraction for the Benchmark Differ (F0).

A ``BenchmarkProfile`` is a comparable, element-typed summary of a document that
can be extracted from EITHER a remediated DOCX (``profile_from_docx``) OR a live
RAWRS ``Document`` (``profile_from_document``), so the two can be diffed with a
single code path. This is deliberately a *structural profile*, not a full
canonical-``Document`` reconstruction (that is a larger effort); it captures only
the element types the F0 gap report measures.

Read-only. No production module imports this.
"""

from __future__ import annotations

import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional
from xml.etree import ElementTree as ET

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
# w:id -1 (separator) and 0 (continuation separator) are structural, not notes.
_SEP_NOTE_IDS = {"-1", "0"}


def _norm(text: Optional[str]) -> str:
    return " ".join((text or "").split()).strip()


@dataclass
class Heading:
    level: int
    text: str


@dataclass
class Table:
    rows: int
    cols: int


@dataclass
class Note:
    marker: str
    text: str


@dataclass
class Figure:
    alt: str


@dataclass
class BenchmarkProfile:
    source: str
    headings: List[Heading] = field(default_factory=list)
    paragraph_count: int = 0
    tables: List[Table] = field(default_factory=list)
    figures: List[Figure] = field(default_factory=list)
    footnotes: List[Note] = field(default_factory=list)
    endnotes: List[Note] = field(default_factory=list)
    page_breaks: int = 0
    printed_page_labels: List[str] = field(default_factory=list)
    # Lines a remediator would treat as running headers/footers: text placed in
    # Word header/footer parts, plus short lines repeated across the body.
    running_header_candidates: List[str] = field(default_factory=list)
    metadata: Dict[str, Optional[str]] = field(default_factory=dict)


def _heading_level(style_name: Optional[str]) -> Optional[int]:
    if not style_name:
        return None
    s = style_name.strip().lower()
    if s == "title":
        return 1
    m = re.match(r"heading\s*([1-9])", s)
    return int(m.group(1)) if m else None


def _level_int(level: Any) -> int:
    """Coerce a RAWRS HeadingLevel (enum, int, or 'H2'/'h2' string) to an int."""
    if isinstance(level, int):
        return level
    val = getattr(level, "value", level)
    if isinstance(val, int):
        return val
    m = re.search(r"([1-9])", str(val))
    return int(m.group(1)) if m else 1


def repeated_short_lines(texts: List[str], min_count: int = 2, max_words: int = 8) -> List[str]:
    """Short lines that repeat — the running-header/masthead/page-number
    signature. Used symmetrically on both sides of the diff."""
    counts = Counter(_norm(t) for t in texts if _norm(t))
    return sorted(t for t, c in counts.items() if c >= min_count and 0 < len(t.split()) <= max_words)


def _notes_from_xml(zf: zipfile.ZipFile, part: str, tag: str) -> List[Note]:
    try:
        root = ET.fromstring(zf.read(part))
    except KeyError:
        return []
    out: List[Note] = []
    for note in root.findall(f"{_W}{tag}"):
        if note.get(f"{_W}id") in _SEP_NOTE_IDS:
            continue
        text = _norm("".join(t.text or "" for t in note.iter(f"{_W}t")))
        if text:
            out.append(Note(marker=note.get(f"{_W}id") or "", text=text))
    return out


def profile_from_docx(path: Path | str) -> BenchmarkProfile:
    """Extract a structural profile from a .docx (benchmark OR RAWRS-generated)."""
    from docx import Document as DocxDocument  # local import: read-only tool dep

    path = Path(path)
    d = DocxDocument(str(path))
    prof = BenchmarkProfile(source=path.name)

    body_texts: List[str] = []
    for p in d.paragraphs:
        text = _norm(p.text)
        if not text:
            continue
        level = _heading_level(getattr(p.style, "name", "") or "")
        if level is not None:
            prof.headings.append(Heading(level=level, text=text))
        else:
            prof.paragraph_count += 1
            body_texts.append(text)

    for t in d.tables:
        rows = len(t.rows)
        cols = len(t.rows[0].cells) if rows else 0
        prof.tables.append(Table(rows=rows, cols=cols))

    prof.metadata = {
        "title": _norm(d.core_properties.title) or None,
        "language": (d.core_properties.language or None),
    }

    header_texts: List[str] = []
    for sec in d.sections:
        for para in list(sec.header.paragraphs) + list(sec.footer.paragraphs):
            if _norm(para.text):
                header_texts.append(_norm(para.text))

    with zipfile.ZipFile(str(path)) as zf:
        prof.footnotes = _notes_from_xml(zf, "word/footnotes.xml", "footnote")
        prof.endnotes = _notes_from_xml(zf, "word/endnotes.xml", "endnote")
        try:
            doc_xml = ET.fromstring(zf.read("word/document.xml"))
            prof.page_breaks = sum(
                1 for br in doc_xml.iter(f"{_W}br") if br.get(f"{_W}type") == "page"
            )
            for docpr in doc_xml.iter(f"{_WP}docPr"):
                prof.figures.append(Figure(alt=_norm(docpr.get("descr"))))
        except KeyError:
            pass

    prof.running_header_candidates = sorted(set(header_texts) | set(repeated_short_lines(body_texts)))
    return prof


def profile_from_document(doc: Any) -> BenchmarkProfile:
    """Extract the same profile shape from a live RAWRS ``Document``.

    Defensive by design (getattr throughout) so it survives model evolution;
    the F0 corpus run compares DOCX-to-DOCX, so this path is for future
    pipeline-integrated measurement.
    """
    prof = BenchmarkProfile(source="rawrs-document")
    body_texts: List[str] = []

    for h in getattr(doc, "headings", []) or []:
        if getattr(h, "is_page_marker", False):
            continue
        prof.headings.append(Heading(level=_level_int(getattr(h, "level", 1)), text=_norm(getattr(h, "text", ""))))

    for b in getattr(doc, "blocks", []) or []:
        text = _norm(getattr(b, "text", ""))
        if text:
            prof.paragraph_count += 1
            body_texts.append(text)

    for tbl in getattr(doc, "tables", []) or []:
        rows_list = getattr(tbl, "rows", []) or []
        cols = len(getattr(rows_list[0], "cells", []) or []) if rows_list else 0
        prof.tables.append(Table(rows=len(rows_list), cols=cols))

    for img in getattr(doc, "images", []) or []:
        fig = getattr(img, "figure", None)
        prof.figures.append(Figure(alt=_norm(getattr(fig, "alt_text", "") if fig is not None else "")))

    for fn in getattr(doc, "footnotes", []) or []:
        note = Note(marker=str(getattr(fn, "number", "") or ""), text=_norm(getattr(fn, "text", "")))
        note_type = getattr(getattr(fn, "note_type", None), "value", None)
        (prof.endnotes if note_type == "endnote" else prof.footnotes).append(note)

    pages = getattr(doc, "pages", []) or []
    prof.page_breaks = max(len(pages) - 1, 0)
    prof.printed_page_labels = [
        str(getattr(p, "printed_label", "")) for p in pages if getattr(p, "printed_label", None)
    ]

    meta = getattr(doc, "metadata", None)
    prof.metadata = {
        "title": _norm(getattr(meta, "title", "")) or None,
        "language": getattr(meta, "language", None),
    }
    prof.running_header_candidates = repeated_short_lines(body_texts)
    return prof
